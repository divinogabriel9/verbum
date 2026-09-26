"""Parish Deck DNA — scaffold, scan, store, and resolve custom Mass masters.

Parishes download a labeled scaffold (content slides only; LiturgyFlow keeps mass
dividers), style fonts/sizes/backgrounds, then upload. Generation clones their
master by slot id (left-footer labels), not hardcoded Theme 1 indices.
"""

from __future__ import annotations

import json
import logging
import re
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping, Optional

logger = logging.getLogger(__name__)

_PROJECT_ROOT = Path(__file__).resolve().parents[1]
_REFERENCE_MASTER = _PROJECT_ROOT / "data" / "reference" / "LFTemplate1.pptx"
_LOCAL_DNA_ROOT = _PROJECT_ROOT / "data" / "parish_dna"
_SCAFFOLD_VERSION = "dna-v1"
_LABEL_PREFIX = "LFDNA"
_LABEL_RE = re.compile(
    rf"^{re.escape(_LABEL_PREFIX)}\|([a-z0-9_]+)\|(REQUIRED|OPTIONAL)\|(\d+)/(\d+)\s*$",
    re.IGNORECASE,
)

# Slots parish media ministries may style. Mass dividers are intentionally absent.
@dataclass(frozen=True)
class DnaSlot:
    key: str
    title: str
    required: bool
    # Extra OPTIONAL copies appended after the stock slide(s) in the scaffold.
    optional_extras: int = 0


DNA_SLOTS: tuple[DnaSlot, ...] = (
    DnaSlot("pre_mass", "Pre-Mass", True),
    DnaSlot("introductory_rites", "Introductory Rites", True),
    DnaSlot("penitential", "Penitential Act", True),
    DnaSlot("kyrie", "Kyrie", True),
    DnaSlot("gloria", "Gloria", True),
    DnaSlot("lotw_prayer", "Opening Prayer (Liturgy of the Word)", False, optional_extras=2),
    DnaSlot("first_reading", "First Reading", True),
    DnaSlot("psalm", "Responsorial Psalm", True),
    DnaSlot("second_reading", "Second Reading", True),
    DnaSlot("gospel_acclamation_alleluia", "Gospel Acclamation · Alleluia", True),
    DnaSlot("gospel_acclamation_dialogue", "Gospel Acclamation · Dialogue", True),
    DnaSlot("gospel_acclamation_end", "Gospel Acclamation · End", True),
    DnaSlot("nicene_creed", "Nicene Creed", True),
    DnaSlot("prayer_faithful", "Prayer of the Faithful", False, optional_extras=2),
    DnaSlot("lote_pray_brethren", "Pray, brethren", True),
    DnaSlot("preface_dialogue", "Preface Dialogue", True),
    DnaSlot("sanctus", "Sanctus", True),
    DnaSlot("mystery_of_faith", "Mystery of Faith", True),
    DnaSlot("great_amen", "Great Amen", True),
    DnaSlot("sign_of_peace", "Sign of Peace", True),
    DnaSlot("lamb_of_god", "Lamb of God", True),
    DnaSlot("communion_rite", "Communion Rite", True),
    DnaSlot("post_communion", "Post-Communion Prayer", False, optional_extras=1),
    DnaSlot("welcoming_newcomers", "Welcoming Newcomers", True),
    DnaSlot("mass_collection", "Mass Collection", True),
    DnaSlot("confession", "Sacrament of Confession", True),
    DnaSlot("final_blessing", "Final Blessing", True),
)

DNA_SLOT_KEYS = frozenset(s.key for s in DNA_SLOTS)
REQUIRED_SLOT_KEYS = frozenset(s.key for s in DNA_SLOTS if s.required)


def format_dna_label(slot_key: str, *, required: bool, part: int, total: int) -> str:
    kind = "REQUIRED" if required else "OPTIONAL"
    return f"{_LABEL_PREFIX}|{slot_key}|{kind}|{part}/{total}"


def parse_dna_label(text: str) -> Optional[dict[str, Any]]:
    raw = (text or "").strip()
    m = _LABEL_RE.match(raw)
    if not m:
        return None
    return {
        "slot": m.group(1).lower(),
        "kind": m.group(2).upper(),
        "part": int(m.group(3)),
        "total": int(m.group(4)),
    }


def _default_master_slide_map() -> dict[str, Any]:
    """Stock Theme 1 indices for DNA slots (from generators.powerpoint._MASTER_SLIDE)."""
    from generators.powerpoint import _MASTER_SLIDE

    out: dict[str, Any] = {}
    for slot in DNA_SLOTS:
        spec = _MASTER_SLIDE.get(slot.key)
        if spec is None:
            continue
        out[slot.key] = list(spec) if isinstance(spec, tuple) else [int(spec)]
    return out


def _stamp_left_footer_label(slide, label: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    from generators.deck_template import SLIDE_HEIGHT, SLIDE_WIDTH

    box = slide.shapes.add_textbox(
        Inches(0.35),
        SLIDE_HEIGHT - Inches(0.42),
        Inches(8.5),
        Inches(0.32),
    )
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = label
    p.alignment = PP_ALIGN.LEFT
    p.font.name = "Consolas"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x88, 0x88, 0x94)


def _copy_slide(dest_prs, src_slide) -> None:
    """Deep-copy a slide's shapes onto a new blank slide (layout-agnostic)."""
    from copy import deepcopy as _deepcopy

    from generators.powerpoint import _layout_blank

    blank = dest_prs.slides.add_slide(_layout_blank(dest_prs))
    # Copy background
    try:
        bg_el = src_slide._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}cSld"
        )
        if bg_el is not None:
            bg = bg_el.find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}bg"
            )
            if bg is not None:
                dest_cSld = blank._element.find(
                    "{http://schemas.openxmlformats.org/presentationml/2006/main}cSld"
                )
                if dest_cSld is not None:
                    existing = dest_cSld.find(
                        "{http://schemas.openxmlformats.org/presentationml/2006/main}bg"
                    )
                    if existing is not None:
                        dest_cSld.remove(existing)
                    dest_cSld.insert(0, _deepcopy(bg))
    except Exception:
        logger.debug("scaffold: could not copy slide background", exc_info=True)

    for shp in src_slide.shapes:
        try:
            newel = _deepcopy(shp.element)
            blank.shapes._spTree.insert_element_before(newel, "p:extLst")
        except Exception:
            logger.debug("scaffold: skip shape copy", exc_info=True)


def build_scaffold_pptx(*, dest: Optional[Path] = None) -> Path:
    """Build a labeled content-only scaffold from LFTemplate1 (no mass dividers)."""
    from pptx import Presentation

    from generators.deck_template import SLIDE_HEIGHT, SLIDE_WIDTH
    from generators.powerpoint import _MASTER_SLIDE, _layout_blank

    if not _REFERENCE_MASTER.is_file():
        raise FileNotFoundError(f"Master template missing: {_REFERENCE_MASTER}")

    src = Presentation(str(_REFERENCE_MASTER))
    out = Presentation()
    out.slide_width = SLIDE_WIDTH
    out.slide_height = SLIDE_HEIGHT

    # Guide cover
    cover = out.slides.add_slide(_layout_blank(out))
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    from generators.powerpoint import _set_slide_bg

    _set_slide_bg(cover, RGBColor(0, 0, 0))
    title_box = cover.shapes.add_textbox(Inches(1), Inches(1.2), Inches(18), Inches(1.2))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "Parish Deck DNA · Scaffold"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xB8, 0x00)
    p.font.name = "Georgia"

    body = cover.shapes.add_textbox(Inches(1), Inches(2.8), Inches(17), Inches(6))
    btf = body.text_frame
    btf.word_wrap = True
    guide_lines = [
        f"Version {_SCAFFOLD_VERSION}",
        "",
        "HOW TO USE",
        "1. Style fonts, sizes, colors, and background photos on content slides.",
        "2. Keep the left-footer LFDNA|… label — do not rename or move it.",
        "3. OPTIONAL slides: delete extras you don’t need, or keep them for multi-page parts",
        "   (e.g. Opening Prayer may be 1–3 slides).",
        "4. Do NOT reorder different sections (First Reading must stay before Psalm, etc.).",
        "5. Mass dividers (LOTW / LOTE posters) are LiturgyFlow-owned — not in this file.",
        "6. Upload the finished .pptx in Settings → Church Profile → Parish Deck DNA.",
        "",
        "REQUIRED slides must remain. OPTIONAL slides may be deleted.",
    ]
    first = True
    for line in guide_lines:
        para = btf.paragraphs[0] if first else btf.add_paragraph()
        first = False
        para.text = line
        para.font.size = Pt(18)
        para.font.color.rgb = RGBColor(0xF0, 0xFD, 0xF4)
        para.font.name = "Arial"
        para.space_after = Pt(4)
    _stamp_left_footer_label(cover, f"{_LABEL_PREFIX}|guide|REQUIRED|1/1")

    stock = _default_master_slide_map()
    for slot in DNA_SLOTS:
        indices = stock.get(slot.key) or []
        if not indices:
            logger.warning("scaffold: no stock slides for slot %s", slot.key)
            continue
        stock_count = len(indices)
        total_parts = stock_count + max(0, int(slot.optional_extras))
        # Stock slides (REQUIRED for required slots; first part of optional slots still kept)
        for i, idx in enumerate(indices):
            if idx < 0 or idx >= len(src.slides):
                continue
            _copy_slide(out, src.slides[idx])
            part = i + 1
            # Stock parts of a required slot are REQUIRED; extras below are OPTIONAL.
            # For optional slots, every part is OPTIONAL so they may delete all.
            is_req = bool(slot.required)
            label = format_dna_label(
                slot.key, required=is_req, part=part, total=total_parts
            )
            _stamp_left_footer_label(out.slides[-1], label)

        # Optional extras: duplicate last stock slide as blank continuation layouts
        donor_idx = indices[-1]
        for extra_i in range(max(0, int(slot.optional_extras))):
            if donor_idx < 0 or donor_idx >= len(src.slides):
                break
            _copy_slide(out, src.slides[donor_idx])
            part = stock_count + extra_i + 1
            label = format_dna_label(
                slot.key, required=False, part=part, total=total_parts
            )
            _stamp_left_footer_label(out.slides[-1], label)

    if dest is None:
        dest = _PROJECT_ROOT / "data" / "reference" / f"ParishDeckDNA_scaffold_{_SCAFFOLD_VERSION}.pptx"
    dest.parent.mkdir(parents=True, exist_ok=True)
    out.save(str(dest))
    return dest.resolve()


def _iter_slide_texts(slide) -> list[str]:
    texts: list[str] = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
            continue
        t = (shape.text_frame.text or "").strip()
        if t:
            texts.append(t)
    return texts


def scan_deck_dna(pptx_path: Path) -> dict[str, Any]:
    """Parse LFDNA labels and build a slide_map. Validates required slots."""
    from pptx import Presentation

    path = Path(pptx_path)
    if not path.is_file():
        return {"ok": False, "error": "File not found.", "slide_map": None}

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        return {"ok": False, "error": f"Could not open PPTX: {exc}", "slide_map": None}

    # slot -> list of (part, slide_index, kind)
    found: dict[str, list[tuple[int, int, str]]] = {}
    unlabeled = 0
    for si, slide in enumerate(prs.slides):
        label_info = None
        for text in _iter_slide_texts(slide):
            # Prefer exact line match; also scan each line
            for line in text.splitlines():
                label_info = parse_dna_label(line.strip())
                if label_info:
                    break
            if label_info:
                break
            # whole-text fallback
            label_info = parse_dna_label(text)
            if label_info:
                break
        if not label_info:
            # Guide slide without parseable multi-part is ok if it says guide
            blob = "\n".join(_iter_slide_texts(slide)).upper()
            if "GUIDE" in blob and si == 0:
                continue
            unlabeled += 1
            continue
        slot = label_info["slot"]
        if slot == "guide":
            continue
        if slot not in DNA_SLOT_KEYS:
            return {
                "ok": False,
                "error": f"Unknown DNA slot on slide {si + 1}: {slot}",
                "slide_map": None,
            }
        found.setdefault(slot, []).append(
            (int(label_info["part"]), si, str(label_info["kind"]))
        )

    slide_map: dict[str, list[int]] = {}
    errors: list[str] = []
    for slot in DNA_SLOTS:
        entries = found.get(slot.key) or []
        if not entries:
            if slot.required:
                errors.append(f"Missing required slot: {slot.key} ({slot.title})")
            continue
        entries.sort(key=lambda t: t[0])
        slide_map[slot.key] = [e[1] for e in entries]

    if errors:
        return {"ok": False, "error": "; ".join(errors), "slide_map": None, "unlabeled": unlabeled}

    return {
        "ok": True,
        "slide_map": slide_map,
        "scaffold_version": _SCAFFOLD_VERSION,
        "slide_count": len(prs.slides),
        "slots_found": sorted(slide_map.keys()),
        "unlabeled": unlabeled,
        "notes": (
            f"Parsed {len(slide_map)} slots across {len(prs.slides)} slides"
            + (f" ({unlabeled} unlabeled skipped)" if unlabeled else "")
            + "."
        ),
    }


def _local_dir(parish_id: str) -> Path:
    safe = re.sub(r"[^a-zA-Z0-9_-]+", "_", (parish_id or "local").strip()) or "local"
    d = _LOCAL_DNA_ROOT / safe
    d.mkdir(parents=True, exist_ok=True)
    return d


def save_dna_local(
    parish_id: str,
    pptx_bytes: bytes,
    slide_map: Mapping[str, Any],
) -> dict[str, Any]:
    """Persist DNA under data/parish_dna/ (dev + offline fallback)."""
    d = _local_dir(parish_id)
    master = d / "master.pptx"
    meta = d / "slide_map.json"
    master.write_bytes(pptx_bytes)
    payload = {
        "scaffold_version": _SCAFFOLD_VERSION,
        "slide_map": dict(slide_map),
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "path": str(master.relative_to(_PROJECT_ROOT)),
    }
    meta.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    try:
        from services.community_store import set_setting

        set_setting("deck_dna_path", str(master))
        set_setting("deck_dna_slide_map", json.dumps(slide_map))
        set_setting("deck_dna_updated_at", payload["updated_at"])
    except Exception:
        logger.debug("local DNA: community_settings write skipped", exc_info=True)
    return {"ok": True, "path": str(master), "slide_map": dict(slide_map), **payload}


def clear_dna_local(parish_id: str) -> dict[str, Any]:
    d = _local_dir(parish_id)
    for name in ("master.pptx", "slide_map.json"):
        p = d / name
        p.unlink(missing_ok=True)
    try:
        from services.community_store import set_setting

        set_setting("deck_dna_path", "")
        set_setting("deck_dna_slide_map", "")
        set_setting("deck_dna_updated_at", "")
    except Exception:
        pass
    return {"ok": True}


def load_dna_local(parish_id: str = "local") -> Optional[dict[str, Any]]:
    d = _local_dir(parish_id)
    master = d / "master.pptx"
    meta = d / "slide_map.json"
    if not master.is_file() or not meta.is_file():
        # Fall back to community_settings (single-tenant local)
        try:
            from services.community_store import get_setting

            path = (get_setting("deck_dna_path") or "").strip()
            raw_map = (get_setting("deck_dna_slide_map") or "").strip()
            if path and Path(path).is_file() and raw_map:
                return {
                    "path": path,
                    "slide_map": json.loads(raw_map),
                    "updated_at": get_setting("deck_dna_updated_at") or "",
                    "scaffold_version": _SCAFFOLD_VERSION,
                }
        except Exception:
            pass
        return None
    try:
        data = json.loads(meta.read_text(encoding="utf-8"))
    except Exception:
        data = {}
    return {
        "path": str(master),
        "slide_map": data.get("slide_map") or {},
        "updated_at": data.get("updated_at") or "",
        "scaffold_version": data.get("scaffold_version") or _SCAFFOLD_VERSION,
    }


def save_dna_supabase(
    parish_id: str,
    pptx_bytes: bytes,
    slide_map: Mapping[str, Any],
) -> dict[str, Any]:
    from services.auth_config import supabase_enabled
    from services.parish_store import get_parish_by_id
    from services.storage_assets import parish_storage_ready, upload_parish_asset

    if not supabase_enabled() or not parish_storage_ready():
        return save_dna_local(parish_id, pptx_bytes, slide_map)

    stored = upload_parish_asset(
        parish_id=parish_id,
        relative_path="deck_dna/master.pptx",
        raw=pptx_bytes,
        content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        upsert=True,
    )
    updated_at = datetime.now(timezone.utc).isoformat()
    try:
        from services.supabase_client import get_service_client

        client = get_service_client()
        client.table("parishes").update(
            {
                "deck_dna_path": stored.path,
                "deck_dna_slide_map": dict(slide_map),
                "deck_dna_updated_at": updated_at,
            }
        ).eq("id", parish_id).execute()
    except Exception as exc:
        logger.warning("DNA: parish column update failed (%s); local mirror kept", exc)
        # Still keep a local materialization for generate
        local = save_dna_local(parish_id, pptx_bytes, slide_map)
        local["storage_path"] = stored.path
        local["warning"] = "Saved to storage; DB columns may need migration."
        return local

    # Materialize local copy for generation (ephemeral FS ok for current process)
    local = save_dna_local(parish_id, pptx_bytes, slide_map)
    return {
        "ok": True,
        "path": local.get("path"),
        "storage_path": stored.path,
        "signed_url": stored.signed_url,
        "slide_map": dict(slide_map),
        "updated_at": updated_at,
        "scaffold_version": _SCAFFOLD_VERSION,
    }


def clear_dna_supabase(parish_id: str) -> dict[str, Any]:
    from services.auth_config import supabase_enabled
    from services.storage_assets import delete_parish_asset, parish_storage_ready

    clear_dna_local(parish_id)
    if supabase_enabled() and parish_storage_ready():
        try:
            delete_parish_asset(parish_id=parish_id, relative_path="deck_dna/master.pptx")
        except Exception:
            logger.debug("DNA: storage delete skipped", exc_info=True)
        try:
            from services.supabase_client import get_service_client

            get_service_client().table("parishes").update(
                {
                    "deck_dna_path": None,
                    "deck_dna_slide_map": None,
                    "deck_dna_updated_at": None,
                }
            ).eq("id", parish_id).execute()
        except Exception:
            logger.debug("DNA: parish clear skipped", exc_info=True)
    return {"ok": True}


def get_dna_status(parish_id: str) -> dict[str, Any]:
    from services.auth_config import supabase_enabled

    pid = (parish_id or "").strip() or "local"
    remote = None
    if supabase_enabled() and pid != "local":
        try:
            from services.parish_store import get_parish_by_id

            parish = get_parish_by_id(pid)
            if parish and parish.get("deck_dna_path"):
                remote = {
                    "storage_path": parish.get("deck_dna_path"),
                    "slide_map": parish.get("deck_dna_slide_map") or {},
                    "updated_at": parish.get("deck_dna_updated_at") or "",
                }
        except Exception:
            logger.debug("DNA status: parish read failed", exc_info=True)

    local = load_dna_local(pid)
    has = bool((remote and remote.get("storage_path")) or (local and local.get("path")))
    return {
        "ok": True,
        "parish_id": pid,
        "has_dna": has,
        "scaffold_version": _SCAFFOLD_VERSION,
        "path": (local or {}).get("path"),
        "storage_path": (remote or {}).get("storage_path") or (local or {}).get("storage_path"),
        "slide_map": (remote or {}).get("slide_map") or (local or {}).get("slide_map") or {},
        "updated_at": (remote or {}).get("updated_at") or (local or {}).get("updated_at") or "",
        "slots": [
            {
                "key": s.key,
                "title": s.title,
                "required": s.required,
                "optional_extras": s.optional_extras,
                "parts": len(((remote or local or {}).get("slide_map") or {}).get(s.key) or []),
            }
            for s in DNA_SLOTS
        ],
    }


def materialize_dna_for_generate(parish_id: str) -> Optional[dict[str, Any]]:
    """Return ``{path, slide_map}`` ready for powerpoint generation, or None."""
    from services.auth_config import supabase_enabled
    from services.storage_assets import download_service_asset, parish_storage_ready

    status = get_dna_status(parish_id)
    if not status.get("has_dna"):
        return None

    slide_map = status.get("slide_map") or {}
    local_path = status.get("path")
    if local_path and Path(local_path).is_file() and slide_map:
        return {"path": local_path, "slide_map": slide_map}

    storage_path = status.get("storage_path")
    if storage_path and supabase_enabled() and parish_storage_ready():
        try:
            raw = download_service_asset(path=storage_path)
            d = _local_dir(parish_id)
            master = d / "master.pptx"
            master.write_bytes(raw)
            if not slide_map:
                scanned = scan_deck_dna(master)
                if scanned.get("ok"):
                    slide_map = scanned["slide_map"]
            return {"path": str(master), "slide_map": slide_map}
        except Exception:
            logger.exception("DNA: failed to materialize from storage")
    return None


def ensure_scaffold_cached() -> Path:
    cached = (
        _PROJECT_ROOT
        / "data"
        / "reference"
        / f"ParishDeckDNA_scaffold_{_SCAFFOLD_VERSION}.pptx"
    )
    if cached.is_file():
        return cached
    return build_scaffold_pptx(dest=cached)
