"""
Multi-slide PowerPoint preview for the web UI.

LibreOffice's ``--convert-to png`` on a ``.pptx`` often emits **only the first slide**
as a single PNG, which makes every preview card look like “slide 1”.

This module prefers: **PPTX → PDF (LibreOffice) → one image per page (pypdfium2)**,
then falls back to whatever PNGs LO produced from a direct PNG export.

Supports progressive slideshow rendering: convert once, emit the first page(s)
immediately, then continue rasterizing the rest into ``out_dir``.
"""

from __future__ import annotations

import logging
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Callable, Optional

logger = logging.getLogger(__name__)

_PROJECT_ROOT = Path(__file__).resolve().parents[1]
_FONTS_DIR = _PROJECT_ROOT / "data" / "reference" / "fonts"

# Preview-only remaps so Mac (system Georgia/Arial) and Linux Present converge.
# Longer names first.
_PREVIEW_FONT_REWRITES: tuple[tuple[bytes, bytes], ...] = (
    (b"Poppins Bold", b"Poppins"),
    (b"Arial Black", b"Arimo"),
    (b"Arial", b"Arimo"),
    (b"Georgia", b"Gelasio"),
    (b"Calibri", b"Carlito"),
)


def count_ppt_slides(ppt: Path) -> int:
    """Return slide count using python-pptx (no LibreOffice)."""
    try:
        from pptx import Presentation
    except ImportError:
        return 0
    try:
        return len(Presentation(str(ppt)).slides)
    except Exception:
        return 0


def _natural_sort_pngs(paths: list[Path]) -> list[Path]:
    """Order ``name1.png``, ``name2.png``, …, ``name10.png`` numerically when possible."""

    def key(p: Path) -> tuple[int, str]:
        m = re.search(r"(\d+)", p.stem)
        return (int(m.group(1)) if m else 0, p.name.lower())

    return sorted(paths, key=key)


def list_slide_images(out_dir: Path) -> list[Path]:
    """Return rendered slide images in order (``slide_0001.jpg`` / ``.png``)."""
    files = [
        p
        for p in out_dir.glob("slide_*")
        if p.is_file() and p.suffix.lower() in {".png", ".jpg", ".jpeg"}
    ]
    return sorted(files, key=lambda p: p.name)


def _rewrite_pptx_fonts_for_preview(src: Path, dest: Path) -> None:
    """Copy PPTX with typeface names remapped to bundled OFL fonts."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(
        dest, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            lower = info.filename.lower()
            if lower.endswith(".xml") or lower.endswith(".rels"):
                for old, new in _PREVIEW_FONT_REWRITES:
                    if old in data:
                        data = data.replace(old, new)
            zout.writestr(info, data)


def _prepare_lo_user_installation() -> tuple[Path, str]:
    """Throwaway LO profile with OFL fonts in ``user/fonts``. Returns (dir, file URI)."""
    profile = Path(tempfile.mkdtemp(prefix="lo-verbum-fonts-"))
    user_fonts = profile / "user" / "fonts"
    user_fonts.mkdir(parents=True, exist_ok=True)
    if _FONTS_DIR.is_dir():
        for ttf in sorted(_FONTS_DIR.glob("*.ttf")):
            try:
                shutil.copy2(ttf, user_fonts / ttf.name)
            except OSError:
                logger.warning("Could not copy preview font %s", ttf.name)
    return profile, profile.resolve().as_uri()


def _write_preview_fontconfig(path: Path) -> None:
    fonts_dir = _FONTS_DIR.resolve()
    alias = ""
    conf_src = _FONTS_DIR / "99-verbum-pptx.conf"
    if conf_src.is_file():
        raw = conf_src.read_text(encoding="utf-8")
        raw = re.sub(r"<\?xml[^>]*\?>", "", raw)
        raw = re.sub(r"<!DOCTYPE[^>]*>", "", raw)
        raw = raw.replace("<fontconfig>", "").replace("</fontconfig>", "")
        alias = raw.strip()
    path.write_text(
        f"""<?xml version="1.0"?>
<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">
<fontconfig>
  <dir>{fonts_dir}</dir>
  <include ignore_missing="yes">/etc/fonts/fonts.conf</include>
  <include ignore_missing="yes">/usr/local/etc/fonts/fonts.conf</include>
  <include ignore_missing="yes">/opt/homebrew/etc/fonts/fonts.conf</include>
  {alias}
  <alias binding="same">
    <family>Georgia</family>
    <prefer><family>Gelasio</family></prefer>
  </alias>
  <alias binding="same">
    <family>Arial</family>
    <prefer><family>Arimo</family></prefer>
  </alias>
  <alias binding="same">
    <family>Calibri</family>
    <prefer><family>Carlito</family></prefer>
  </alias>
</fontconfig>
""",
        encoding="utf-8",
    )


def convert_pptx_to_pdf(
    ppt: Path,
    out_dir: Path,
    *,
    soffice_bin: str,
    timeout: int = 180,
) -> Optional[Path]:
    """Convert ``ppt`` to a PDF in ``out_dir`` via LibreOffice. Returns the PDF path or ``None``."""
    out_dir.mkdir(parents=True, exist_ok=True)
    profile_dir: Optional[Path] = None
    tmp_dir: Optional[Path] = None
    fc_dir: Optional[Path] = None
    work_ppt = ppt
    try:
        profile_dir, profile_uri = _prepare_lo_user_installation()
        tmp_dir = Path(tempfile.mkdtemp(prefix="ppt-preview-"))
        tmp_ppt = tmp_dir / f"{ppt.stem}_preview.pptx"
        try:
            _rewrite_pptx_fonts_for_preview(ppt, tmp_ppt)
            work_ppt = tmp_ppt
        except Exception:
            logger.warning("Preview font rewrite failed; converting original PPTX", exc_info=True)
            work_ppt = ppt

        env = os.environ.copy()
        try:
            fc_dir = Path(tempfile.mkdtemp(prefix="fc-verbum-"))
            fc_file = fc_dir / "fonts.conf"
            _write_preview_fontconfig(fc_file)
            env["FONTCONFIG_FILE"] = str(fc_file.resolve())
        except Exception:
            logger.warning("Could not write preview fontconfig", exc_info=True)

        cmd = [
            soffice_bin,
            f"-env:UserInstallation={profile_uri}",
            "--headless",
            "--norestore",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(work_ppt.resolve()),
        ]
        try:
            subprocess.run(
                cmd, check=True, capture_output=True, text=True, timeout=timeout, env=env
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
            return None

        for stem in (work_ppt.stem, ppt.stem):
            candidate = out_dir / f"{stem}.pdf"
            if candidate.is_file():
                final = out_dir / f"{ppt.stem}.pdf"
                if candidate.resolve() != final.resolve():
                    try:
                        if final.exists():
                            final.unlink()
                        candidate.rename(final)
                        return final
                    except OSError:
                        return candidate
                return candidate
        pdfs = sorted(out_dir.glob("*.pdf"))
        return pdfs[0] if pdfs else None
    finally:
        if tmp_dir and tmp_dir.exists():
            shutil.rmtree(tmp_dir, ignore_errors=True)
        if profile_dir and profile_dir.exists():
            shutil.rmtree(profile_dir, ignore_errors=True)
        if fc_dir and fc_dir.exists():
            shutil.rmtree(fc_dir, ignore_errors=True)


def _save_pil_image(pil_image, dest: Path, *, image_format: str) -> bool:
    fmt = (image_format or "png").strip().lower()
    try:
        if fmt in {"jpg", "jpeg"}:
            if dest.suffix.lower() not in {".jpg", ".jpeg"}:
                dest = dest.with_suffix(".jpg")
            rgb = pil_image.convert("RGB") if pil_image.mode != "RGB" else pil_image
            rgb.save(dest, format="JPEG", quality=90, optimize=False)
        else:
            if dest.suffix.lower() != ".png":
                dest = dest.with_suffix(".png")
            # optimize=False — much faster for large decks; size difference is minor.
            pil_image.save(dest, format="PNG", optimize=False)
        return True
    except OSError:
        return False


def _slide_dest(out_dir: Path, index_1based: int, image_format: str) -> Path:
    fmt = (image_format or "png").strip().lower()
    ext = ".jpg" if fmt in {"jpg", "jpeg"} else ".png"
    return out_dir / f"slide_{index_1based:04d}{ext}"


def render_pdf_page_range(
    pdf_path: Path,
    out_dir: Path,
    *,
    start: int = 0,
    end: Optional[int] = None,
    scale: float = 1.25,
    image_format: str = "png",
    on_page: Optional[Callable[[int, Path], None]] = None,
) -> list[Path]:
    """
    Rasterize PDF pages ``start..end`` (end exclusive; None = through last page).

    ``start``/``end`` are 0-based page indices.
    """
    try:
        import pypdfium2 as pdfium  # type: ignore[import-untyped]
    except ImportError:
        return []

    out_dir.mkdir(parents=True, exist_ok=True)
    out: list[Path] = []
    try:
        doc = pdfium.PdfDocument(str(pdf_path))
    except Exception:
        return []

    render_scale = max(0.5, float(scale or 1.25))
    n = len(doc)
    stop = n if end is None else min(int(end), n)
    begin = max(0, int(start))
    try:
        for i in range(begin, stop):
            page = doc[i]
            dest = _slide_dest(out_dir, i + 1, image_format)
            try:
                bitmap = page.render(scale=render_scale)
                pil_image = bitmap.to_pil()
                if not _save_pil_image(pil_image, dest, image_format=image_format):
                    continue
                out.append(dest)
                if on_page:
                    try:
                        on_page(i + 1, dest)
                    except Exception:
                        pass
            except Exception:
                continue
            finally:
                try:
                    page.close()
                except Exception:
                    pass
    finally:
        try:
            doc.close()
        except Exception:
            pass
    return out


def count_pdf_pages(pdf_path: Path) -> int:
    try:
        import pypdfium2 as pdfium  # type: ignore[import-untyped]
    except ImportError:
        return 0
    try:
        doc = pdfium.PdfDocument(str(pdf_path))
    except Exception:
        return 0
    try:
        return len(doc)
    finally:
        try:
            doc.close()
        except Exception:
            pass


def begin_progressive_ppt_preview(
    ppt: Path,
    out_dir: Path,
    *,
    soffice_bin: str,
    scale: float = 1.25,
    image_format: str = "jpeg",
    first_batch: int = 2,
) -> tuple[list[Path], Optional[Path], int, str]:
    """
    Convert PPTX→PDF into ``out_dir``, render the first ``first_batch`` pages.

    Returns ``(first_paths, pdf_path, total_pages, message)``.
    Caller should keep rendering via ``render_pdf_page_range`` starting at ``len(first_paths)``.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    for p in out_dir.glob("*"):
        if p.is_file():
            p.unlink(missing_ok=True)

    n_slides = count_ppt_slides(ppt)
    pdf_path = convert_pptx_to_pdf(ppt, out_dir, soffice_bin=soffice_bin)
    if pdf_path is None:
        return ([], None, n_slides, "LibreOffice could not convert the deck to PDF.")

    # Normalize name so status/background always find it.
    stable_pdf = out_dir / "_deck.pdf"
    if pdf_path.resolve() != stable_pdf.resolve():
        try:
            if stable_pdf.exists():
                stable_pdf.unlink(missing_ok=True)
            pdf_path.replace(stable_pdf)
            pdf_path = stable_pdf
        except OSError:
            pdf_path = pdf_path  # keep LO name

    total = count_pdf_pages(pdf_path) or n_slides
    batch = max(1, int(first_batch or 1))
    first = render_pdf_page_range(
        pdf_path,
        out_dir,
        start=0,
        end=min(batch, total or batch),
        scale=scale,
        image_format=image_format,
    )
    msg = ""
    if n_slides and total and n_slides != total:
        msg = (
            f"Preview: PDF has {total} page(s) (python-pptx reports {n_slides} slides)."
        )
    return (first, pdf_path, total or len(first), msg)


def render_ppt_preview_pngs(
    ppt: Path,
    out_dir: Path,
    *,
    soffice_bin: str,
    scale: float = 1.25,
    image_format: str = "png",
) -> tuple[list[Path], str]:
    """
    Write one image per slide into ``out_dir`` when possible.

    Returns ``(paths_sorted, message)``. ``message`` is non-empty when a fallback was used.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    n_slides = count_ppt_slides(ppt)
    message = ""

    # --- Primary: PDF export then pypdfium2 (full deck) ---
    with tempfile.TemporaryDirectory(prefix="ppt_pdf_") as tmp:
        tmp_path = Path(tmp)
        pdf_path = convert_pptx_to_pdf(ppt, tmp_path, soffice_bin=soffice_bin)

        if pdf_path is not None:
            # Clear prior slide images only (keep other files if any).
            for old in list_slide_images(out_dir):
                old.unlink(missing_ok=True)
            pdf_pages = render_pdf_page_range(
                pdf_path, out_dir, start=0, end=None, scale=scale, image_format=image_format
            )
            if pdf_pages:
                if n_slides and len(pdf_pages) != n_slides:
                    message = (
                        f"Preview: {len(pdf_pages)} image(s) from PDF (python-pptx reports {n_slides} slides). "
                        "Counts can differ if the deck has hidden slides; install current LibreOffice + pypdfium2."
                    )
                return (sorted(pdf_pages, key=lambda p: p.name), message)

    # --- Fallback: direct PNG export (often first slide only) ---
    png_cmd = [
        soffice_bin,
        "--headless",
        "--convert-to",
        "png",
        "--outdir",
        str(out_dir),
        str(ppt.resolve()),
    ]
    try:
        subprocess.run(png_cmd, check=True, capture_output=True, text=True, timeout=120)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
        return ([], "LibreOffice could not render the deck.")

    pngs = _natural_sort_pngs([p for p in out_dir.glob("*.png") if p.is_file()])
    if len(pngs) == 1 and n_slides > 1:
        message = (
            "Preview shows the first slide only (LibreOffice PNG export). "
            "Install ``pypdfium2`` (`pip install pypdfium2`) for full-deck thumbnails via PDF."
        )
    return (pngs, message)
