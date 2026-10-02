#!/usr/bin/env python3
"""Build an external partner introduction deck for LiturgyFlow.

Audience: Catholic music / publishing partners.
Content is client/partner-facing only — no technical internals.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "presentations"
OUT_PATH = OUT_DIR / "LiturgyFlow_Publisher_Introduction.pptx"
ICON_PATH = ROOT / "static" / "brand" / "app-icon.png"

# Brand (public design tokens — berry primary, calm ink)
BERRY = RGBColor(0xA1, 0x0F, 0x0D)
INK = RGBColor(0x15, 0x33, 0x3D)
MUTED = RGBColor(0x5C, 0x6B, 0x75)
LINE = RGBColor(0xD8, 0xDE, 0xE2)
SOFT = RGBColor(0xF6, 0xF3, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x16, 0x65, 0x34)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _set_run_font(run, *, size_pt: float, bold: bool = False, color: RGBColor = INK, name: str = "Calibri") -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _set_para(p, text: str, *, size: float, bold: bool = False, color: RGBColor = INK, align=PP_ALIGN.LEFT, space_after: float = 6, space_before: float = 0, name: str = "Calibri"):
    p.clear()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    _set_run_font(run, size_pt=size, bold=bold, color=color, name=name)
    return p


def _fill_solid(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _add_rect(slide, left, top, width, height, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _fill_solid(shape, color)
    return shape


def _add_round_rect(slide, left, top, width, height, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _fill_solid(shape, color)
    # Softer corners
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def _set_cell_border(cell, color: RGBColor = LINE, width_pt: float = 0.75) -> None:
    hex_val = "%02X%02X%02X" % (color[0], color[1], color[2])
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        ln = OxmlElement(f"a:{edge}")
        ln.set("w", str(int(width_pt * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", hex_val)
        sf.append(srgb)
        ln.append(sf)
        prst = OxmlElement("a:prstDash")
        prst.set("val", "solid")
        ln.append(prst)
        tcPr.append(ln)


def _style_table_cell(cell, text: str, *, bold: bool = False, size: float = 14, color: RGBColor = INK, fill: RGBColor | None = None, align=PP_ALIGN.LEFT):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.background()
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _set_para(p, text, size=size, bold=bold, color=color, align=align, space_after=0)
    _set_cell_border(cell)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def _footer(slide, page: int, total: int = 11) -> None:
    # Left accent line
    _add_rect(slide, Inches(0.7), Inches(7.05), Inches(11.9), Emu(9525), LINE)
    box = _add_textbox(slide, Inches(0.7), Inches(7.1), Inches(8), Inches(0.3))
    _set_para(box.text_frame.paragraphs[0], "LiturgyFlow  ·  Confidential partner introduction", size=10, color=MUTED, space_after=0)
    num = _add_textbox(slide, Inches(11.2), Inches(7.1), Inches(1.4), Inches(0.3))
    _set_para(num.text_frame.paragraphs[0], f"{page} / {total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT, space_after=0)


def _section_kicker(slide, text: str, top=Inches(0.55)) -> None:
    box = _add_textbox(slide, Inches(0.7), top, Inches(11.5), Inches(0.35))
    _set_para(box.text_frame.paragraphs[0], text.upper(), size=11, bold=True, color=BERRY, space_after=0, name="Calibri")


def _title(slide, text: str, top=Inches(0.85)) -> None:
    box = _add_textbox(slide, Inches(0.7), top, Inches(11.8), Inches(0.7))
    _set_para(box.text_frame.paragraphs[0], text, size=28, bold=True, color=INK, space_after=0, name="Calibri")


def _new_slide(prs: Presentation):
    blank = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(blank)
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    # Left brand rail
    _add_rect(slide, 0, 0, Inches(0.12), SLIDE_H, BERRY)
    return slide


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ——— 1. Title ———
    s = _new_slide(prs)
    _add_rect(s, 0, 0, SLIDE_W, SLIDE_H, SOFT)
    _add_rect(s, 0, 0, Inches(0.18), SLIDE_H, BERRY)
    if ICON_PATH.is_file():
        s.shapes.add_picture(str(ICON_PATH), Inches(0.85), Inches(1.55), height=Inches(0.85))
    title = _add_textbox(s, Inches(0.85), Inches(2.55), Inches(11), Inches(0.9))
    _set_para(title.text_frame.paragraphs[0], "LiturgyFlow", size=44, bold=True, color=INK, space_after=0)
    sub = _add_textbox(s, Inches(0.85), Inches(3.4), Inches(11), Inches(1.1))
    tf = sub.text_frame
    tf.word_wrap = True
    _set_para(
        tf.paragraphs[0],
        "A simple web-based tool that helps Catholic church media teams\nprepare Mass presentations and related materials with less repetitive manual work.",
        size=18,
        color=MUTED,
        space_after=12,
    )
    line = _add_textbox(s, Inches(0.85), Inches(5.0), Inches(11), Inches(0.8))
    _set_para(line.text_frame.paragraphs[0], "Introduction for music & publishing partners", size=14, bold=True, color=BERRY, space_after=4)
    _set_para(line.text_frame.add_paragraph(), "Confidential  ·  For discussion purposes", size=12, color=MUTED, space_after=0)
    _footer(s, 1)

    # ——— 2. Problem ———
    s = _new_slide(prs)
    _section_kicker(s, "Context")
    _title(s, "The weekly Mass presentation problem")
    lead = _add_textbox(s, Inches(0.7), Inches(1.65), Inches(11.8), Inches(0.6))
    _set_para(
        lead.text_frame.paragraphs[0],
        "Every week, parish media and music teams typically need to prepare:",
        size=16,
        color=MUTED,
        space_after=0,
    )
    cards = [
        ("Projection slides", "Readings, responses, and hymn lyrics for screen display"),
        ("Printed materials", "A Mass leaflet when paper copies are needed"),
        ("Supporting visuals", "Materials that help the assembly follow the liturgy"),
    ]
    for i, (h, d) in enumerate(cards):
        left = Inches(0.7 + i * 4.05)
        _add_round_rect(s, left, Inches(2.45), Inches(3.85), Inches(2.15), SOFT)
        _add_rect(s, left, Inches(2.45), Inches(0.12), Inches(2.15), BERRY)
        hb = _add_textbox(s, left + Inches(0.35), Inches(2.7), Inches(3.3), Inches(0.45))
        _set_para(hb.text_frame.paragraphs[0], h, size=16, bold=True, color=INK, space_after=0)
        db = _add_textbox(s, left + Inches(0.35), Inches(3.25), Inches(3.3), Inches(1.1))
        db.text_frame.word_wrap = True
        _set_para(db.text_frame.paragraphs[0], d, size=14, color=MUTED, space_after=0)
    close = _add_textbox(s, Inches(0.7), Inches(5.0), Inches(11.8), Inches(1.2))
    tf = close.text_frame
    tf.word_wrap = True
    _set_para(tf.paragraphs[0], "The challenge", size=14, bold=True, color=BERRY, space_after=6)
    for t in (
        "The work is recurring and deadline-driven.",
        "Content comes from multiple places — lectionary, Order of Mass, hymnals, and parish preferences.",
        "Manual assembly takes time that music and media volunteers often do not have.",
    ):
        p = tf.add_paragraph()
        run = p.add_run()
        run.text = f"•  {t}"
        _set_run_font(run, size_pt=14, color=INK)
        p.space_after = Pt(4)
    _footer(s, 2)

    # ——— 3. Traditional workflow ———
    s = _new_slide(prs)
    _section_kicker(s, "Today")
    _title(s, "The traditional preparation workflow")
    steps = [
        "Look up the Sunday’s readings and liturgical texts",
        "Copy or retype content into PowerPoint",
        "Format slides for readability on screen",
        "Look up hymn lyrics from hymnals or files",
        "Create hymn lyric slides verse by verse",
        "Adjust layout, fonts, and parish branding",
        "Optionally prepare a printed leaflet or divider materials",
        "Review, correct, and finalize before Mass",
    ]
    for i, step in enumerate(steps):
        col = i % 2
        row = i // 2
        left = Inches(0.7 + col * 6.2)
        top = Inches(1.75 + row * 1.05)
        _add_round_rect(s, left, top, Inches(5.95), Inches(0.9), SOFT)
        num = _add_textbox(s, left + Inches(0.2), top + Inches(0.22), Inches(0.55), Inches(0.45))
        _set_para(num.text_frame.paragraphs[0], f"{i + 1:02d}", size=16, bold=True, color=BERRY, space_after=0)
        tx = _add_textbox(s, left + Inches(0.75), top + Inches(0.25), Inches(5.0), Inches(0.5))
        tx.text_frame.word_wrap = True
        _set_para(tx.text_frame.paragraphs[0], step, size=13, color=INK, space_after=0)
    _footer(s, 3)

    # ——— 4. How LiturgyFlow helps ———
    s = _new_slide(prs)
    _section_kicker(s, "Approach")
    _title(s, "How LiturgyFlow simplifies the workflow")
    lead = _add_textbox(s, Inches(0.7), Inches(1.65), Inches(11.8), Inches(0.7))
    lead.text_frame.word_wrap = True
    _set_para(
        lead.text_frame.paragraphs[0],
        "A web-based preparation workspace that helps parish teams move from “this Sunday’s Mass” to ready-to-use materials with fewer manual steps.",
        size=16,
        color=MUTED,
        space_after=0,
    )
    left_items = [
        "Select the Mass date / celebration",
        "Review liturgical content for that Mass",
        "Choose the hymns for the celebration",
        "Generate Mass presentation & related materials",
        "Review the output before use",
    ]
    right_items = [
        "Repeated copying and reformatting",
        "Rebuilding the same slide structure each week",
        "Manual lyric-slide assembly for selected hymns",
    ]
    _add_round_rect(s, Inches(0.7), Inches(2.5), Inches(5.9), Inches(3.5), SOFT)
    _add_round_rect(s, Inches(6.85), Inches(2.5), Inches(5.75), Inches(3.5), SOFT)
    lh = _add_textbox(s, Inches(0.95), Inches(2.7), Inches(5.4), Inches(0.4))
    _set_para(lh.text_frame.paragraphs[0], "What the team does", size=14, bold=True, color=BERRY, space_after=0)
    rh = _add_textbox(s, Inches(7.1), Inches(2.7), Inches(5.3), Inches(0.4))
    _set_para(rh.text_frame.paragraphs[0], "What this reduces", size=14, bold=True, color=BERRY, space_after=0)
    lb = _add_textbox(s, Inches(0.95), Inches(3.2), Inches(5.4), Inches(2.5))
    lb.text_frame.word_wrap = True
    for i, t in enumerate(left_items):
        p = lb.text_frame.paragraphs[0] if i == 0 else lb.text_frame.add_paragraph()
        run = p.add_run()
        run.text = f"•  {t}"
        _set_run_font(run, size_pt=14, color=INK)
        p.space_after = Pt(8)
    rb = _add_textbox(s, Inches(7.1), Inches(3.2), Inches(5.3), Inches(2.5))
    rb.text_frame.word_wrap = True
    for i, t in enumerate(right_items):
        p = rb.text_frame.paragraphs[0] if i == 0 else rb.text_frame.add_paragraph()
        run = p.add_run()
        run.text = f"•  {t}"
        _set_run_font(run, size_pt=14, color=INK)
        p.space_after = Pt(8)
    note = _add_textbox(s, Inches(0.7), Inches(6.2), Inches(11.8), Inches(0.45))
    _set_para(
        note.text_frame.paragraphs[0],
        "The goal is not to replace pastoral judgment — it is to reduce repetitive production work.",
        size=13,
        bold=True,
        color=INK,
        space_after=0,
    )
    _footer(s, 4)

    # ——— 5. Outputs ———
    s = _new_slide(prs)
    _section_kicker(s, "Outputs")
    _title(s, "What LiturgyFlow can produce")
    outputs = [
        ("Mass PowerPoint", "A complete presentation for projection, including liturgical texts and hymn lyric slides as selected by the parish team."),
        ("A4 Mass leaflet", "A printable half-fold leaflet for offline or paper use when copies are needed."),
        ("Mass divider / poster", "Visual materials suited to marking parts of the Mass or supporting related parish presentation needs."),
        ("Related media", "Additional presentation and media materials that support the same Mass preparation workflow."),
    ]
    for i, (h, d) in enumerate(outputs):
        top = Inches(1.7 + i * 1.15)
        _add_round_rect(s, Inches(0.7), top, Inches(11.9), Inches(1.05), SOFT)
        _add_rect(s, Inches(0.7), top, Inches(0.12), Inches(1.05), BERRY if i < 3 else GREEN)
        hb = _add_textbox(s, Inches(1.1), top + Inches(0.15), Inches(11.2), Inches(0.35))
        _set_para(hb.text_frame.paragraphs[0], h, size=16, bold=True, color=INK, space_after=0)
        db = _add_textbox(s, Inches(1.1), top + Inches(0.5), Inches(11.2), Inches(0.45))
        db.text_frame.word_wrap = True
        _set_para(db.text_frame.paragraphs[0], d, size=13, color=MUTED, space_after=0)
    foot = _add_textbox(s, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.4))
    foot.text_frame.word_wrap = True
    _set_para(
        foot.text_frame.paragraphs[0],
        "Parishes remain responsible for appropriate permissions for copyrighted content they display or distribute.",
        size=11,
        color=MUTED,
        space_after=0,
    )
    _footer(s, 5)

    # ——— 6. Sunday workflow ———
    s = _new_slide(prs)
    _section_kicker(s, "Example")
    _title(s, "Example Sunday Mass workflow")
    flow = [
        ("1", "Select Mass", "Choose the Sunday or celebration"),
        ("2", "Prepare & review", "Confirm readings and liturgical texts"),
        ("3", "Select hymns", "Assign hymns for the celebration"),
        ("4", "Generate", "Produce PowerPoint and optional leaflet"),
        ("5", "Review", "Check slides and print materials"),
        ("6", "Present", "Use the finished materials at Mass"),
    ]
    for i, (n, h, d) in enumerate(flow):
        left = Inches(0.55 + i * 2.1)
        _add_round_rect(s, left, Inches(2.2), Inches(1.95), Inches(3.2), SOFT)
        nb = _add_textbox(s, left + Inches(0.15), Inches(2.45), Inches(1.65), Inches(0.55))
        _set_para(nb.text_frame.paragraphs[0], n, size=28, bold=True, color=BERRY, align=PP_ALIGN.CENTER, space_after=0)
        hb = _add_textbox(s, left + Inches(0.12), Inches(3.15), Inches(1.7), Inches(0.7))
        hb.text_frame.word_wrap = True
        _set_para(hb.text_frame.paragraphs[0], h, size=13, bold=True, color=INK, align=PP_ALIGN.CENTER, space_after=0)
        db = _add_textbox(s, left + Inches(0.12), Inches(3.9), Inches(1.7), Inches(1.2))
        db.text_frame.word_wrap = True
        _set_para(db.text_frame.paragraphs[0], d, size=11, color=MUTED, align=PP_ALIGN.CENTER, space_after=0)
        if i < 5:
            arrow = _add_textbox(s, left + Inches(1.85), Inches(3.4), Inches(0.3), Inches(0.4))
            _set_para(arrow.text_frame.paragraphs[0], "→", size=16, bold=True, color=BERRY, align=PP_ALIGN.CENTER, space_after=0)
    close = _add_textbox(s, Inches(0.7), Inches(5.75), Inches(11.8), Inches(0.5))
    _set_para(
        close.text_frame.paragraphs[0],
        "One continuous preparation path — from Mass selection to presentation.",
        size=15,
        bold=True,
        color=INK,
        space_after=0,
    )
    _footer(s, 6)

    # ——— 7. Hymn workflow ———
    s = _new_slide(prs)
    _section_kicker(s, "Hymns")
    _title(s, "Hymn workflow within Mass preparation")
    points = [
        "Hymn content is stored in LiturgyFlow because it is required to generate a complete Mass presentation.",
        "When a hymn is selected for a Mass, the corresponding lyrics can be automatically formatted into the appropriate PowerPoint slides.",
        "This allows music and media teams to focus on choosing suitable music for the liturgy, rather than rebuilding lyric slides by hand each week.",
        "Hymn selection remains a pastoral and music-ministry decision. LiturgyFlow supports formatting and inclusion within the Mass presentation workflow.",
    ]
    for i, t in enumerate(points):
        top = Inches(1.75 + i * 1.1)
        _add_round_rect(s, Inches(0.7), top, Inches(11.9), Inches(0.95), SOFT)
        nb = _add_textbox(s, Inches(0.95), top + Inches(0.25), Inches(0.55), Inches(0.45))
        _set_para(nb.text_frame.paragraphs[0], f"{i + 1}", size=18, bold=True, color=BERRY, space_after=0)
        tx = _add_textbox(s, Inches(1.55), top + Inches(0.22), Inches(10.8), Inches(0.6))
        tx.text_frame.word_wrap = True
        _set_para(tx.text_frame.paragraphs[0], t, size=14, color=INK, space_after=0)
    note = _add_textbox(s, Inches(0.7), Inches(6.25), Inches(11.9), Inches(0.4))
    _set_para(
        note.text_frame.paragraphs[0],
        "Hymn content exists in the product to serve Mass preparation — not as a separate lyrics destination.",
        size=13,
        bold=True,
        color=MUTED,
        space_after=0,
    )
    _footer(s, 7)

    # ——— 8. Copyright ———
    s = _new_slide(prs)
    _section_kicker(s, "Rights")
    _title(s, "Hymn copyright & publisher relationship")
    stmts = [
        "LiturgyFlow is not intended to operate as a standalone hymn-lyrics publishing service.",
        "Hymn content is used only to complete Mass presentation generation — not as a separate lyrics product.",
        "LiturgyFlow does not claim ownership of copyrighted songs or lyrics belonging to third parties.",
        "We seek appropriate permission from rights holders for applicable repertoire used in this workflow.",
        "Parishes remain responsible for securing the rights required for their own use, display, and distribution.",
    ]
    for i, t in enumerate(stmts):
        top = Inches(1.7 + i * 0.85)
        _add_round_rect(s, Inches(0.7), top, Inches(11.9), Inches(0.75), SOFT)
        _add_rect(s, Inches(0.7), top, Inches(0.12), Inches(0.75), BERRY)
        tx = _add_textbox(s, Inches(1.1), top + Inches(0.18), Inches(11.2), Inches(0.45))
        tx.text_frame.word_wrap = True
        _set_para(tx.text_frame.paragraphs[0], t, size=14, color=INK, space_after=0)
    close = _add_textbox(s, Inches(0.7), Inches(6.15), Inches(11.9), Inches(0.5))
    close.text_frame.word_wrap = True
    _set_para(
        close.text_frame.paragraphs[0],
        "We approach hymn repertoire with respect for publishers, composers, and the Church’s obligation to honor intellectual property.",
        size=13,
        bold=True,
        color=INK,
        space_after=0,
    )
    _footer(s, 8)

    # ——— 9. Free vs Paid ———
    s = _new_slide(prs)
    _section_kicker(s, "Public pricing")
    _title(s, "Free tier and paid plans")
    lead = _add_textbox(s, Inches(0.7), Inches(1.6), Inches(11.8), Inches(0.45))
    lead.text_frame.word_wrap = True
    _set_para(
        lead.text_frame.paragraphs[0],
        "Parishes can start on a free tier. Paid plans unlock fuller preparation for regular weekly use.",
        size=15,
        color=MUTED,
        space_after=0,
    )

    # Free card
    _add_round_rect(s, Inches(0.7), Inches(2.25), Inches(5.9), Inches(4.0), SOFT)
    _add_rect(s, Inches(0.7), Inches(2.25), Inches(5.9), Inches(0.55), MUTED)
    fh = _add_textbox(s, Inches(0.95), Inches(2.35), Inches(5.4), Inches(0.4))
    _set_para(fh.text_frame.paragraphs[0], "Free tier", size=18, bold=True, color=WHITE, space_after=0)
    free_items = [
        ("Included", "Mass presentation generation within a limited allowance"),
        ("Limited", "Limited number of generations"),
        ("Not included", "Custom parish PowerPoint template"),
        ("Not included", "Beautifully crafted Mass divider posters"),
    ]
    for i, (label, text) in enumerate(free_items):
        top = Inches(3.0 + i * 0.75)
        lb = _add_textbox(s, Inches(0.95), top, Inches(5.4), Inches(0.28))
        _set_para(lb.text_frame.paragraphs[0], label.upper(), size=11, bold=True, color=BERRY if label != "Included" else GREEN, space_after=0)
        tb = _add_textbox(s, Inches(0.95), top + Inches(0.26), Inches(5.4), Inches(0.4))
        tb.text_frame.word_wrap = True
        _set_para(tb.text_frame.paragraphs[0], text, size=14, color=INK, space_after=0)

    # Paid card
    _add_round_rect(s, Inches(6.85), Inches(2.25), Inches(5.75), Inches(4.0), SOFT)
    _add_rect(s, Inches(6.85), Inches(2.25), Inches(5.75), Inches(0.55), BERRY)
    ph = _add_textbox(s, Inches(7.1), Inches(2.35), Inches(5.3), Inches(0.4))
    _set_para(ph.text_frame.paragraphs[0], "Paid plans", size=18, bold=True, color=WHITE, space_after=0)
    paid_items = [
        ("Included", "Fuller generation capacity for regular weekly Mass preparation"),
        ("Included", "Custom parish PowerPoint template"),
        ("Included", "Beautifully crafted Mass divider posters"),
        ("Included", "Materials suited to ongoing parish media ministry"),
    ]
    for i, (label, text) in enumerate(paid_items):
        top = Inches(3.0 + i * 0.75)
        lb = _add_textbox(s, Inches(7.1), top, Inches(5.3), Inches(0.28))
        _set_para(lb.text_frame.paragraphs[0], label.upper(), size=11, bold=True, color=GREEN, space_after=0)
        tb = _add_textbox(s, Inches(7.1), top + Inches(0.26), Inches(5.3), Inches(0.4))
        tb.text_frame.word_wrap = True
        _set_para(tb.text_frame.paragraphs[0], text, size=14, color=INK, space_after=0)
    _footer(s, 9)

    # ——— 10. Pricing table ———
    s = _new_slide(prs)
    _section_kicker(s, "Public pricing")
    _title(s, "Paid subscription pricing")
    lead = _add_textbox(s, Inches(0.7), Inches(1.6), Inches(11.8), Inches(0.4))
    _set_para(
        lead.text_frame.paragraphs[0],
        "Simple public subscription plans for parish use (shown in local currencies):",
        size=15,
        color=MUTED,
        space_after=0,
    )
    rows = [
        ("Plan", "KRW", "PHP", "MYR", "USD"),
        ("Monthly", "₩9,900", "₱199", "RM19.90", "$6.99"),
        ("3-month", "₩27,000", "₱549", "RM54.90", "$18.99"),
        ("6-month", "₩49,000", "₱999", "RM99.90", "$34.99"),
        ("Annual", "₩79,000", "₱1,599", "RM159.90", "$59.99"),
    ]
    table = s.shapes.add_table(len(rows), 5, Inches(0.7), Inches(2.2), Inches(11.9), Inches(3.4)).table
    widths = [Inches(2.4), Inches(2.35), Inches(2.35), Inches(2.4), Inches(2.4)]
    for i, w in enumerate(widths):
        table.columns[i].width = w
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            is_header = r == 0
            is_first_col = c == 0
            _style_table_cell(
                table.cell(r, c),
                val,
                bold=is_header or is_first_col,
                size=15 if is_header else 14,
                color=WHITE if is_header else INK,
                fill=BERRY if is_header else (SOFT if r % 2 == 0 else WHITE),
                align=PP_ALIGN.CENTER if c > 0 or is_header else PP_ALIGN.LEFT,
            )
    note = _add_textbox(s, Inches(0.7), Inches(5.9), Inches(11.9), Inches(0.7))
    note.text_frame.word_wrap = True
    _set_para(
        note.text_frame.paragraphs[0],
        "Pricing shown is the public parish subscription offering. A free tier is also available with limited generation.",
        size=13,
        color=INK,
        space_after=4,
    )
    _set_para(
        note.text_frame.add_paragraph(),
        "Publisher permissions for hymn repertoire are separate from parish subscription pricing.",
        size=13,
        color=MUTED,
        space_after=0,
    )
    _footer(s, 10)

    # ——— 11. Collaboration ———
    s = _new_slide(prs)
    _section_kicker(s, "Request")
    _title(s, "What we are asking of publishers")
    lead = _add_textbox(s, Inches(0.7), Inches(1.6), Inches(11.8), Inches(0.9))
    lead.text_frame.word_wrap = True
    _set_para(
        lead.text_frame.paragraphs[0],
        "We are seeking permission to include applicable hymn repertoire so that parish Mass presentation generation can be complete — not to operate a separate hymn-lyrics service.",
        size=16,
        color=MUTED,
        space_after=0,
    )
    goals = [
        ("Not a lyrics product", "Hymns are not sold, browsed, or offered as a standalone publishing feature."),
        ("Only for Mass decks", "Content is used so selected hymns can become properly formatted slides in the Mass presentation."),
        ("Permission to complete", "We ask for permission so the Mass-generation workflow can include your repertoire responsibly."),
    ]
    for i, (h, d) in enumerate(goals):
        left = Inches(0.7 + i * 4.05)
        _add_round_rect(s, left, Inches(2.65), Inches(3.85), Inches(2.15), SOFT)
        _add_rect(s, left, Inches(2.65), Inches(0.12), Inches(2.15), BERRY)
        hb = _add_textbox(s, left + Inches(0.35), Inches(2.85), Inches(3.3), Inches(0.45))
        hb.text_frame.word_wrap = True
        _set_para(hb.text_frame.paragraphs[0], h, size=14, bold=True, color=INK, space_after=0)
        db = _add_textbox(s, left + Inches(0.35), Inches(3.35), Inches(3.3), Inches(1.2))
        db.text_frame.word_wrap = True
        _set_para(db.text_frame.paragraphs[0], d, size=13, color=MUTED, space_after=0)
    discuss = _add_textbox(s, Inches(0.7), Inches(5.05), Inches(11.8), Inches(1.15))
    discuss.text_frame.word_wrap = True
    _set_para(discuss.text_frame.paragraphs[0], "What permission would enable", size=14, bold=True, color=BERRY, space_after=6)
    for t in (
        "Parish teams can select approved hymns and generate complete Mass PowerPoint decks, including lyric slides",
        "Rights holders retain ownership; LiturgyFlow claims no ownership of third-party songs or lyrics",
        "Attribution or other acknowledgment can be included where you prefer",
    ):
        p = discuss.text_frame.add_paragraph()
        run = p.add_run()
        run.text = f"•  {t}"
        _set_run_font(run, size_pt=13, color=INK)
        p.space_after = Pt(3)
    close = _add_textbox(s, Inches(0.7), Inches(6.3), Inches(11.9), Inches(0.4))
    close.text_frame.word_wrap = True
    _set_para(
        close.text_frame.paragraphs[0],
        "We welcome a conversation about permission that helps parishes prepare Mass materials while respecting your rights.",
        size=13,
        bold=True,
        color=INK,
        space_after=0,
    )
    _footer(s, 11)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PATH))
    return OUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
