#!/usr/bin/env python3
"""Build DRAFT MOA deliverables for music/publishing partner discussions.

Outputs:
  docs/agreements/LiturgyFlow_Publisher_MOA_Draft.docx
  docs/agreements/LiturgyFlow_Publisher_MOA_Draft.pptx
  docs/agreements/LiturgyFlow_Publisher_MOA_Draft.pdf
  docs/agreements/LiturgyFlow_Publisher_MOA_Draft_Presentation.pdf

Content mirrors docs/agreements/LiturgyFlow_Publisher_MOA_Draft.md.
Discussion draft only — no invented commercial/legal terms.

PDF export uses LibreOffice (soffice) when available.
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches as PptInches, Pt as PptPt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "agreements"
DOCX_PATH = OUT_DIR / "LiturgyFlow_Publisher_MOA_Draft.docx"
PPTX_PATH = OUT_DIR / "LiturgyFlow_Publisher_MOA_Draft.pptx"
PDF_DOC_PATH = OUT_DIR / "LiturgyFlow_Publisher_MOA_Draft.pdf"
PDF_PPT_PATH = OUT_DIR / "LiturgyFlow_Publisher_MOA_Draft_Presentation.pdf"
ICON_PATH = ROOT / "static" / "brand" / "app-icon.png"

# Brand
BERRY = PptRGB(0xA1, 0x0F, 0x0D)
INK = PptRGB(0x15, 0x33, 0x3D)
MUTED = PptRGB(0x5C, 0x6B, 0x75)
LINE = PptRGB(0xD8, 0xDE, 0xE2)
SOFT = PptRGB(0xF6, 0xF3, 0xF0)
WHITE = PptRGB(0xFF, 0xFF, 0xFF)

DOC_BERRY = RGBColor(0xA1, 0x0F, 0x0D)
DOC_INK = RGBColor(0x15, 0x33, 0x3D)
DOC_MUTED = RGBColor(0x5C, 0x6B, 0x75)

SLIDE_W = PptInches(13.333)
SLIDE_H = PptInches(7.5)
TOTAL_SLIDES = 12

TBA = "[To be agreed by the parties]"

LICENSE_TOPICS = [
    ("Repertoire (titles, catalogs, editions, languages)", TBA),
    ("Permitted uses and any restrictions", TBA),
    ("Territory", TBA),
    ("Duration / term of license", TBA),
    ("Exclusivity or non-exclusivity", TBA),
    ("Sublicensing (if any) to churches or affiliates", TBA),
    ("Delivery, updates, and withdrawal of titles", TBA),
    ("Attribution / copyright notice requirements", TBA),
    ("Other commercial or usage conditions", TBA),
]


# ── DOCX helpers ─────────────────────────────────────────────────────────────

def _set_run(run, *, size: float = 11, bold: bool = False, italic: bool = False, color: RGBColor = DOC_INK, name: str = "Calibri"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def _p(doc: Document, text: str = "", *, size: float = 11, bold: bool = False, italic: bool = False, color: RGBColor = DOC_INK, align=WD_ALIGN_PARAGRAPH.LEFT, space_after: float = 8, space_before: float = 0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if text:
        run = p.add_run(text)
        _set_run(run, size=size, bold=bold, italic=italic, color=color)
    return p


def _rich(doc: Document, parts: list[tuple[str, dict]], *, align=WD_ALIGN_PARAGRAPH.LEFT, space_after: float = 8, space_before: float = 0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for text, style in parts:
        run = p.add_run(text)
        _set_run(run, **style)
    return p


def _heading(doc: Document, text: str):
    return _p(doc, text, size=13, bold=True, color=DOC_BERRY, space_before=14, space_after=6)


def _bullet(doc: Document, text: str, *, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        _set_run(r1, size=11, bold=True, color=DOC_INK)
        r2 = p.add_run(text)
        _set_run(r2, size=11, color=DOC_INK)
    else:
        run = p.add_run(text)
        _set_run(run, size=11, color=DOC_INK)
    return p


def _sig_line(doc: Document, label: str):
    _p(doc, label, size=11, color=DOC_INK, space_after=2, space_before=2)


def build_docx() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = DOC_INK

    _p(doc, "DRAFT — FOR DISCUSSION PURPOSES ONLY", size=14, bold=True, color=DOC_BERRY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    _p(doc, "MEMORANDUM OF AGREEMENT", size=18, bold=True, color=DOC_INK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _p(
        doc,
        "Concerning Potential Collaboration and Licensing of Hymn Content",
        size=12,
        bold=True,
        color=DOC_MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=10,
    )
    _p(
        doc,
        "This document is an initial business discussion draft and is subject to review and revision by the parties and their respective legal advisers. It is not a final legal agreement and does not create binding commercial, licensing, or financial obligations unless and until the parties execute a definitive written agreement.",
        size=10,
        italic=True,
        color=DOC_MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )
    _p(doc, f"Date of draft: {TBA}", size=10, color=DOC_MUTED, space_after=2)
    _p(doc, "Reference: LiturgyFlow × Jesuit Music Ministry (JMM) — Hymn Content Collaboration (Draft MOA)", size=10, color=DOC_MUTED, space_after=14)

    # 1
    _heading(doc, "1. PARTIES")
    _p(doc, "This Memorandum of Agreement (“MOA”) is entered into for discussion purposes by and between:")
    _p(doc, "LiturgyFlow", bold=True, space_after=2)
    _p(doc, "(“LiturgyFlow” or the “Platform Operator”)", size=10, color=DOC_MUTED, space_after=2)
    _p(doc, f"[Legal entity name, address, and authorized representative — {TBA}]", size=10, italic=True, color=DOC_MUTED, space_after=8)
    _p(doc, "and", space_after=8)
    _p(doc, "Jesuit Music Ministry (JMM)", bold=True, space_after=2)
    _p(doc, "(“Publisher” or “JMM”)", size=10, color=DOC_MUTED, space_after=2)
    _p(doc, f"[Legal entity name, address, and authorized representative — {TBA}]", size=10, italic=True, color=DOC_MUTED, space_after=8)
    _p(doc, "LiturgyFlow and JMM are referred to individually as a “Party” and collectively as the “Parties.”")

    # 2
    _heading(doc, "2. PURPOSE")
    _p(
        doc,
        "The Parties wish to explore a collaboration under which applicable hymn and related musical content controlled or administered by the Publisher may be used within LiturgyFlow’s Mass presentation-generation workflow for Catholic churches and parish communities.",
    )
    _p(
        doc,
        "This MOA sets out the Parties’ shared understanding of the intended collaboration, the nature of LiturgyFlow’s use of hymn content, and the principal topics that would need to be agreed in any subsequent definitive agreement. It is intended to facilitate constructive discussion and does not itself constitute a completed license, partnership, joint venture, or revenue-sharing arrangement.",
    )

    # 3
    _heading(doc, "3. LITURGYFLOW")
    _p(
        doc,
        "LiturgyFlow is a web-based tool that helps Catholic churches and parish communities prepare Mass presentations and related liturgical media. Among other capabilities, LiturgyFlow assists users in assembling materials commonly used in the celebration of Mass, including readings and hymn-related presentation slides.",
    )
    _p(
        doc,
        f"LiturgyFlow serves churches, parish media teams, and related Catholic communities who prepare presentation materials for worship. Further operational and commercial details of the Platform Operator’s business are {TBA} for inclusion in any definitive agreement, as appropriate.",
    )

    # 4
    _heading(doc, "4. HYMNS AND CONTENT")
    _p(
        doc,
        "For LiturgyFlow to generate complete Mass presentation materials, applicable hymn content—including lyrics—must be available within the platform when a church selects a hymn for inclusion in a Mass presentation.",
    )
    _p(
        doc,
        "Accordingly, where licensed and authorized, LiturgyFlow may store applicable hymn content from the Publisher’s repertoire for the limited purpose of enabling presentation generation. When a church selects a hymn, LiturgyFlow may automatically format the corresponding lyrics into the resulting PowerPoint (or similar) Mass presentation.",
    )
    _p(
        doc,
        f"Nothing in this section is intended to expand the scope of permitted use beyond what the Parties later agree in writing. The specific repertoire, formats, metadata, delivery methods, and update processes for such content are {TBA}.",
    )

    # 5
    _heading(doc, "5. INTENDED USE")
    _p(
        doc,
        "The intended use of Publisher hymn content under any collaboration contemplated by this MOA is the incorporation of that content into church Mass presentation materials prepared through LiturgyFlow (for example, lyric slides within a Mass PowerPoint deck used by a parish in connection with worship).",
    )
    _p(doc, "LiturgyFlow is not intended to operate as:", space_after=4)
    _bullet(doc, "an independent hymn-lyrics publishing service; or")
    _bullet(doc, "a standalone lyrics-distribution service for general public browsing, download, or redistribution of lyrics outside the Mass presentation-generation workflow.")
    _p(
        doc,
        f"Any additional or different uses (including public display, print reproduction, streaming, recording, or redistribution outside Mass presentation materials) are outside the scope of this draft and are {TBA}, if at all.",
        space_before=6,
    )

    # 6
    _heading(doc, "6. COPYRIGHT OWNERSHIP")
    _p(doc, "Each Party retains all right, title, and interest in and to its respective intellectual property.")
    _p(doc, "Without limiting the foregoing:", space_after=4)
    _bullet(doc, "LiturgyFlow does not claim ownership of third-party copyrighted songs, lyrics, musical compositions, or other Publisher-controlled content.")
    _bullet(doc, "The Publisher does not, by discussing this collaboration, claim ownership of LiturgyFlow’s platform, trademarks, branding, software, workflows, or other LiturgyFlow intellectual property.")
    _bullet(doc, f"Ownership of any jointly developed materials, if any, is {TBA}.")
    _p(doc, "Any license contemplated herein is a permission to use specified content on agreed terms, not a transfer of ownership.", space_before=6)

    # 7
    _heading(doc, "7. LICENSE")
    _p(
        doc,
        "Subject to execution of a definitive agreement, the Publisher may grant LiturgyFlow an appropriate permission and/or license to use the agreed repertoire of hymn content within LiturgyFlow for the intended uses described in Section 5.",
    )
    _p(doc, "The following license terms remain open and are expressly not decided by this draft:", space_after=6)

    table = doc.add_table(rows=1 + len(LICENSE_TOPICS), cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = ""
    hdr[1].text = ""
    hdr[0].paragraphs[0].clear()
    hdr[1].paragraphs[0].clear()
    r0 = hdr[0].paragraphs[0].add_run("Topic")
    _set_run(r0, size=10, bold=True, color=DOC_INK)
    r1 = hdr[1].paragraphs[0].add_run("Status")
    _set_run(r1, size=10, bold=True, color=DOC_INK)
    for i, (topic, status) in enumerate(LICENSE_TOPICS, start=1):
        c0, c1 = table.rows[i].cells
        c0.paragraphs[0].clear()
        c1.paragraphs[0].clear()
        t_run = c0.paragraphs[0].add_run(topic)
        _set_run(t_run, size=10, color=DOC_INK)
        s_run = c1.paragraphs[0].add_run(status)
        _set_run(s_run, size=10, italic=True, color=DOC_MUTED)
    _p(doc, "No license is granted by this MOA alone.", space_before=10)

    # 8
    _heading(doc, "8. COMMERCIAL ARRANGEMENT")
    _p(
        doc,
        "The Parties welcome a good-faith conversation about an appropriate commercial arrangement associated with permission to include applicable Publisher repertoire within LiturgyFlow’s Mass presentation workflow — so that parish teams can prepare complete Mass materials while respecting the Publisher’s rights.",
    )
    _p(
        doc,
        "Any such arrangement would be informed by factors such as:",
        space_after=4,
    )
    for item in (
        "the applicable repertoire;",
        "the intended use within LiturgyFlow’s Mass presentation workflow;",
        "territory;",
        "the number and type of users (for example, churches, parishes, or other authorized accounts); and",
        "the usage model.",
    ):
        _bullet(doc, item)
    _p(
        doc,
        "Publisher permission arrangements are separate from LiturgyFlow’s public parish subscription pricing.",
        space_before=6,
    )
    _rich(
        doc,
        [
            ("Specific commercial terms — including any licensing fees, royalties, or other consideration, if applicable — are ", {"size": 11, "color": DOC_INK}),
            (TBA, {"size": 11, "bold": True, "color": DOC_BERRY}),
            (" and would be set out only in a definitive written agreement (or schedules thereto).", {"size": 11, "color": DOC_INK}),
        ],
    )
    _p(doc, "This MOA does not establish any fee, royalty, or payment obligation.")

    # 9
    _heading(doc, "9. REPORTING")
    _p(
        doc,
        f"The Parties contemplate that LiturgyFlow may provide reasonable usage reports concerning the licensed repertoire (for example, information helpful to understanding which titles are selected or used within the Mass presentation workflow), in a form and frequency that is {TBA}.",
    )
    _p(
        doc,
        "Reporting under any definitive agreement is intended to support transparency regarding use of the licensed repertoire. It is not intended to require disclosure of LiturgyFlow’s complete financial records, overall business finances, or unrelated commercial information.",
    )
    _rich(
        doc,
        [
            ("Any financial reporting beyond agreed usage reporting, and any audit rights or audit procedures, are ", {"size": 11, "color": DOC_INK}),
            (TBA, {"size": 11, "bold": True, "color": DOC_BERRY}),
            (" and, if included, would be limited to what is reasonably necessary for the licensed repertoire and related fees.", {"size": 11, "color": DOC_INK}),
        ],
    )

    # 10
    _heading(doc, "10. CONFIDENTIALITY")
    _p(
        doc,
        "Each Party agrees to treat as confidential any non-public business, technical, and commercial information exchanged by the other Party in connection with discussions under this MOA, and to use such information only for evaluating and negotiating the contemplated collaboration.",
    )
    _p(
        doc,
        "Confidential information does not include information that: (a) is or becomes publicly available other than through breach of this section; (b) was already known to the receiving Party without confidentiality obligation; (c) is independently developed without use of the other Party’s confidential information; or (d) is rightfully received from a third party without confidentiality restriction.",
    )
    _p(
        doc,
        f"Nothing in this section requires LiturgyFlow to disclose proprietary technical details of its platform beyond what LiturgyFlow chooses to share for discussion purposes. Duration of confidentiality obligations, permitted disclosures (including to legal advisers), and related remedies are {TBA}.",
    )

    # 11
    _heading(doc, "11. TERM AND TERMINATION")
    _rich(
        doc,
        [
            ("The duration of any definitive agreement, renewal terms, and termination rights (including termination for convenience, for cause, and effect of termination on licensed content and materials already generated) are ", {"size": 11, "color": DOC_INK}),
            (TBA, {"size": 11, "bold": True, "color": DOC_BERRY}),
            (".", {"size": 11, "color": DOC_INK}),
        ],
    )
    _p(
        doc,
        "Until a definitive agreement is signed, either Party may discontinue discussions at any time. Discontinuation of discussions under this draft MOA does not, by itself, create liability for the other Party, except as may arise under separately agreed confidentiality obligations, if any.",
    )

    # 12
    _heading(doc, "12. LIMITATIONS")
    _p(doc, "Any collaboration or license arising from these discussions would apply only to:", space_after=4)
    _bullet(doc, "the repertoire expressly agreed by the Parties; and")
    _bullet(doc, "the uses expressly agreed by the Parties.")
    _p(
        doc,
        "No broader rights are implied. Content not included in the agreed repertoire, and uses not expressly authorized, remain outside the scope of any contemplated arrangement.",
        space_before=6,
    )
    _rich(
        doc,
        [
            ("Governing law, dispute resolution, indemnity, limitation of liability, warranties, and other legal terms customarily addressed in a definitive agreement are ", {"size": 11, "color": DOC_INK}),
            (TBA, {"size": 11, "bold": True, "color": DOC_BERRY}),
            (" and are not established by this draft.", {"size": 11, "color": DOC_INK}),
        ],
    )

    # 13
    _heading(doc, "13. SIGNATURES")
    _p(
        doc,
        "By signing below, the Parties acknowledge that they have reviewed this draft for discussion purposes. Signatures on this draft MOA confirm willingness to continue good-faith discussions and do not, by themselves, create a binding license or commercial payment obligation, unless the Parties expressly state otherwise in a subsequent definitive written agreement.",
        space_after=14,
    )

    _p(doc, "FOR LITURGYFLOW", bold=True, space_after=10)
    for label in ("Authorized signature: _______________________________", "Name: _____________________________________________", "Title: _____________________________________________", "Date: _____________________________________________"):
        _sig_line(doc, label)
    _p(doc, "", space_after=12)
    _p(doc, "FOR JESUIT MUSIC MINISTRY (JMM)", bold=True, space_after=10)
    for label in ("Authorized signature: _______________________________", "Name: _____________________________________________", "Title: _____________________________________________", "Date: _____________________________________________"):
        _sig_line(doc, label)

    _p(doc, "End of Draft MOA — For Discussion Purposes Only", size=10, italic=True, color=DOC_MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18)

    doc.save(DOCX_PATH)
    return DOCX_PATH


# ── PPTX helpers ─────────────────────────────────────────────────────────────

def _set_run_font(run, *, size_pt: float, bold: bool = False, color: PptRGB = INK, name: str = "Calibri") -> None:
    run.font.name = name
    run.font.size = PptPt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _set_para(p, text: str, *, size: float, bold: bool = False, color: PptRGB = INK, align=PP_ALIGN.LEFT, space_after: float = 6, space_before: float = 0, name: str = "Calibri"):
    p.clear()
    p.alignment = align
    p.space_after = PptPt(space_after)
    p.space_before = PptPt(space_before)
    run = p.add_run()
    run.text = text
    _set_run_font(run, size_pt=size, bold=bold, color=color, name=name)
    return p


def _fill_solid(shape, color: PptRGB) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _add_rect(slide, left, top, width, height, color: PptRGB):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _fill_solid(shape, color)
    return shape


def _add_round_rect(slide, left, top, width, height, color: PptRGB):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _fill_solid(shape, color)
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def _footer(slide, page: int) -> None:
    _add_rect(slide, PptInches(0.7), PptInches(7.05), PptInches(11.9), Emu(9525), LINE)
    box = _add_textbox(slide, PptInches(0.7), PptInches(7.1), PptInches(8.5), PptInches(0.3))
    _set_para(box.text_frame.paragraphs[0], "LiturgyFlow  ·  Draft MOA  ·  For discussion only", size=10, color=MUTED, space_after=0)
    num = _add_textbox(slide, PptInches(11.2), PptInches(7.1), PptInches(1.4), PptInches(0.3))
    _set_para(num.text_frame.paragraphs[0], f"{page} / {TOTAL_SLIDES}", size=10, color=MUTED, align=PP_ALIGN.RIGHT, space_after=0)


def _section_kicker(slide, text: str, top=PptInches(0.45)) -> None:
    box = _add_textbox(slide, PptInches(0.7), top, PptInches(11.5), PptInches(0.3))
    _set_para(box.text_frame.paragraphs[0], text.upper(), size=11, bold=True, color=BERRY, space_after=0)


def _title(slide, text: str, top=PptInches(0.75)) -> None:
    box = _add_textbox(slide, PptInches(0.7), top, PptInches(11.8), PptInches(0.55))
    _set_para(box.text_frame.paragraphs[0], text, size=26, bold=True, color=INK, space_after=0)


def _new_slide(prs: Presentation):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    _add_rect(slide, 0, 0, PptInches(0.12), SLIDE_H, BERRY)
    return slide


def _bullets(slide, left, top, width, height, items: list[str], *, size: float = 14):
    box = _add_textbox(slide, left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"•  {t}"
        _set_run_font(run, size_pt=size, color=INK)
        p.space_after = PptPt(8)
    return box


def _set_cell_border(cell, color: PptRGB = LINE, width_pt: float = 0.75) -> None:
    hex_val = "%02X%02X%02X" % (color[0], color[1], color[2])
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        ln = OxmlElement(f"a:{edge}")
        ln.set("w", str(int(width_pt * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", hex_val)
        sf.append(srgb)
        ln.append(sf)
        prst = OxmlElement("a:prstDash")
        prst.set("val", "solid")
        ln.append(prst)
        tcPr.append(ln)


def _style_table_cell(cell, text: str, *, bold: bool = False, size: float = 12, color: PptRGB = INK, fill: PptRGB | None = None):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.background()
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _set_para(p, text, size=size, bold=bold, color=color, space_after=0)
    _set_cell_border(cell)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def build_pptx() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1 Title
    s = _new_slide(prs)
    _add_rect(s, 0, 0, SLIDE_W, SLIDE_H, SOFT)
    _add_rect(s, 0, 0, PptInches(0.18), SLIDE_H, BERRY)
    if ICON_PATH.is_file():
        s.shapes.add_picture(str(ICON_PATH), PptInches(0.85), PptInches(1.35), height=PptInches(0.75))
    banner = _add_textbox(s, PptInches(0.85), PptInches(2.25), PptInches(11), PptInches(0.4))
    _set_para(banner.text_frame.paragraphs[0], "DRAFT — FOR DISCUSSION PURPOSES ONLY", size=14, bold=True, color=BERRY, space_after=0)
    title = _add_textbox(s, PptInches(0.85), PptInches(2.7), PptInches(11), PptInches(0.7))
    _set_para(title.text_frame.paragraphs[0], "Memorandum of Agreement", size=36, bold=True, color=INK, space_after=0)
    sub = _add_textbox(s, PptInches(0.85), PptInches(3.45), PptInches(11), PptInches(0.9))
    sub.text_frame.word_wrap = True
    _set_para(
        sub.text_frame.paragraphs[0],
        "Potential collaboration and licensing of hymn content\nbetween LiturgyFlow and Jesuit Music Ministry (JMM)",
        size=16,
        color=MUTED,
        space_after=0,
    )
    note = _add_textbox(s, PptInches(0.85), PptInches(4.7), PptInches(11), PptInches(1.2))
    note.text_frame.word_wrap = True
    _set_para(
        note.text_frame.paragraphs[0],
        "This document is an initial business discussion draft and is subject to review and revision by the parties and their respective legal advisers.",
        size=13,
        color=INK,
        space_after=8,
    )
    p = note.text_frame.add_paragraph()
    run = p.add_run()
    run.text = "Open commercial and legal terms are marked:  [To be agreed by the parties]"
    _set_run_font(run, size_pt=13, bold=True, color=BERRY)
    _footer(s, 1)

    # 2 Parties & Purpose
    s = _new_slide(prs)
    _section_kicker(s, "Sections 1–2")
    _title(s, "Parties and purpose")
    _add_round_rect(s, PptInches(0.7), PptInches(1.55), PptInches(5.85), PptInches(2.35), SOFT)
    _add_round_rect(s, PptInches(6.75), PptInches(1.55), PptInches(5.85), PptInches(2.35), SOFT)
    h1 = _add_textbox(s, PptInches(0.95), PptInches(1.7), PptInches(5.4), PptInches(0.35))
    _set_para(h1.text_frame.paragraphs[0], "LiturgyFlow", size=16, bold=True, color=BERRY, space_after=0)
    b1 = _add_textbox(s, PptInches(0.95), PptInches(2.15), PptInches(5.4), PptInches(1.5))
    b1.text_frame.word_wrap = True
    _set_para(b1.text_frame.paragraphs[0], "Platform Operator\nWeb-based Catholic Mass presentation tool", size=14, color=INK, space_after=8)
    p = b1.text_frame.add_paragraph()
    run = p.add_run()
    run.text = f"Legal entity details: {TBA}"
    _set_run_font(run, size_pt=12, color=MUTED)
    h2 = _add_textbox(s, PptInches(7.0), PptInches(1.7), PptInches(5.4), PptInches(0.35))
    _set_para(h2.text_frame.paragraphs[0], "Jesuit Music Ministry (JMM)", size=16, bold=True, color=BERRY, space_after=0)
    b2 = _add_textbox(s, PptInches(7.0), PptInches(2.15), PptInches(5.4), PptInches(1.5))
    b2.text_frame.word_wrap = True
    _set_para(b2.text_frame.paragraphs[0], "Publisher (“JMM”)\nCatholic music ministry / publishing partner", size=14, color=INK, space_after=8)
    p = b2.text_frame.add_paragraph()
    run = p.add_run()
    run.text = f"Legal entity details: {TBA}"
    _set_run_font(run, size_pt=12, color=MUTED)
    purpose = _add_textbox(s, PptInches(0.7), PptInches(4.2), PptInches(11.9), PptInches(2.2))
    purpose.text_frame.word_wrap = True
    _set_para(purpose.text_frame.paragraphs[0], "Purpose", size=14, bold=True, color=BERRY, space_after=6)
    _set_para(
        purpose.text_frame.add_paragraph(),
        "Explore collaboration so applicable Publisher hymn content may be used within LiturgyFlow’s Mass presentation-generation workflow for Catholic churches and parish communities.",
        size=15,
        color=INK,
        space_after=8,
    )
    _set_para(
        purpose.text_frame.add_paragraph(),
        "This draft facilitates discussion. It is not a completed license, partnership, joint venture, or revenue-sharing arrangement.",
        size=14,
        color=MUTED,
        space_after=0,
    )
    _footer(s, 2)

    # 3 LiturgyFlow
    s = _new_slide(prs)
    _section_kicker(s, "Section 3")
    _title(s, "What LiturgyFlow is")
    lead = _add_textbox(s, PptInches(0.7), PptInches(1.55), PptInches(11.9), PptInches(1.1))
    lead.text_frame.word_wrap = True
    _set_para(
        lead.text_frame.paragraphs[0],
        "A web-based tool that helps Catholic churches and parish communities prepare Mass presentations and related liturgical media.",
        size=18,
        color=INK,
        space_after=0,
    )
    cards = [
        ("For churches", "Parish media and music teams preparing worship materials"),
        ("Mass materials", "Presentations that may include readings and hymn lyric slides"),
        ("Partner context", "Discussion focuses on hymn content use — not platform internals"),
    ]
    for i, (h, d) in enumerate(cards):
        left = PptInches(0.7 + i * 4.05)
        _add_round_rect(s, left, PptInches(3.0), PptInches(3.85), PptInches(2.4), SOFT)
        _add_rect(s, left, PptInches(3.0), PptInches(0.12), PptInches(2.4), BERRY)
        hb = _add_textbox(s, left + PptInches(0.35), PptInches(3.25), PptInches(3.3), PptInches(0.45))
        _set_para(hb.text_frame.paragraphs[0], h, size=15, bold=True, color=INK, space_after=0)
        db = _add_textbox(s, left + PptInches(0.35), PptInches(3.85), PptInches(3.3), PptInches(1.3))
        db.text_frame.word_wrap = True
        _set_para(db.text_frame.paragraphs[0], d, size=14, color=MUTED, space_after=0)
    _footer(s, 3)

    # 4 Hymns & content
    s = _new_slide(prs)
    _section_kicker(s, "Section 4")
    _title(s, "Hymns and content")
    _bullets(
        s,
        PptInches(0.7),
        PptInches(1.6),
        PptInches(11.9),
        PptInches(3.2),
        [
            "Complete Mass presentations require applicable hymn content, including lyrics, when a church selects a hymn.",
            "Where licensed and authorized, LiturgyFlow may store applicable Publisher hymn content for presentation generation.",
            "When a hymn is selected, LiturgyFlow may automatically format the lyrics into the Mass PowerPoint (or similar) presentation.",
            f"Specific repertoire, formats, metadata, delivery, and update processes: {TBA}.",
        ],
        size=16,
    )
    note = _add_textbox(s, PptInches(0.7), PptInches(5.3), PptInches(11.9), PptInches(1.0))
    note.text_frame.word_wrap = True
    _set_para(
        note.text_frame.paragraphs[0],
        "Scope of permitted use is limited to what the Parties later agree in writing.",
        size=14,
        bold=True,
        color=BERRY,
        space_after=0,
    )
    _footer(s, 4)

    # 5 Intended use
    s = _new_slide(prs)
    _section_kicker(s, "Section 5")
    _title(s, "Intended use")
    _add_round_rect(s, PptInches(0.7), PptInches(1.55), PptInches(11.9), PptInches(1.6), SOFT)
    yes = _add_textbox(s, PptInches(0.95), PptInches(1.75), PptInches(11.4), PptInches(1.2))
    yes.text_frame.word_wrap = True
    _set_para(yes.text_frame.paragraphs[0], "Intended", size=13, bold=True, color=BERRY, space_after=4)
    _set_para(
        yes.text_frame.add_paragraph(),
        "Incorporation of Publisher hymn content into church Mass presentation materials prepared through LiturgyFlow (for example, lyric slides in a parish Mass deck).",
        size=15,
        color=INK,
        space_after=0,
    )
    no_h = _add_textbox(s, PptInches(0.7), PptInches(3.5), PptInches(11.9), PptInches(0.4))
    _set_para(no_h.text_frame.paragraphs[0], "Not intended", size=13, bold=True, color=BERRY, space_after=0)
    _bullets(
        s,
        PptInches(0.7),
        PptInches(3.95),
        PptInches(11.9),
        PptInches(2.2),
        [
            "An independent hymn-lyrics publishing service",
            "A standalone lyrics-distribution service for general public browsing, download, or redistribution outside the Mass presentation workflow",
            f"Any additional uses (print, streaming, recording, etc.): {TBA}, if at all",
        ],
        size=15,
    )
    _footer(s, 5)

    # 6 Copyright
    s = _new_slide(prs)
    _section_kicker(s, "Section 6")
    _title(s, "Copyright ownership")
    lead = _add_textbox(s, PptInches(0.7), PptInches(1.55), PptInches(11.9), PptInches(0.6))
    _set_para(lead.text_frame.paragraphs[0], "Each Party retains its respective intellectual property rights.", size=18, bold=True, color=INK, space_after=0)
    _bullets(
        s,
        PptInches(0.7),
        PptInches(2.4),
        PptInches(11.9),
        PptInches(3.5),
        [
            "LiturgyFlow does not claim ownership of third-party copyrighted songs or lyrics.",
            "The Publisher does not, by discussing this collaboration, claim ownership of LiturgyFlow’s platform, trademarks, branding, software, or workflows.",
            f"Ownership of any jointly developed materials, if any: {TBA}.",
            "Any license is a permission to use specified content on agreed terms — not a transfer of ownership.",
        ],
        size=16,
    )
    _footer(s, 6)

    # 7 License
    s = _new_slide(prs)
    _section_kicker(s, "Section 7")
    _title(s, "License — terms not decided in this draft")
    lead = _add_textbox(s, PptInches(0.7), PptInches(1.4), PptInches(11.9), PptInches(0.7))
    lead.text_frame.word_wrap = True
    _set_para(
        lead.text_frame.paragraphs[0],
        "Subject to a definitive agreement, the Publisher may grant LiturgyFlow an appropriate permission/license for the agreed repertoire. No license is granted by this MOA alone.",
        size=14,
        color=MUTED,
        space_after=0,
    )
    rows = min(5, len(LICENSE_TOPICS))
    table = s.shapes.add_table(1 + rows, 2, PptInches(0.7), PptInches(2.2), PptInches(11.9), PptInches(3.6)).table
    table.columns[0].width = PptInches(7.4)
    table.columns[1].width = PptInches(4.5)
    _style_table_cell(table.cell(0, 0), "Topic", bold=True, size=12, color=INK, fill=SOFT)
    _style_table_cell(table.cell(0, 1), "Status", bold=True, size=12, color=INK, fill=SOFT)
    for i, (topic, status) in enumerate(LICENSE_TOPICS[:rows]):
        _style_table_cell(table.cell(i + 1, 0), topic, size=12, color=INK)
        _style_table_cell(table.cell(i + 1, 1), status, size=11, color=MUTED)
    more = _add_textbox(s, PptInches(0.7), PptInches(6.0), PptInches(11.9), PptInches(0.5))
    _set_para(
        more.text_frame.paragraphs[0],
        "Also open: exclusivity, sublicensing, delivery/updates, attribution, and other commercial conditions — all [To be agreed by the parties].",
        size=12,
        color=MUTED,
        space_after=0,
    )
    _footer(s, 7)

    # 8 Commercial arrangement
    s = _new_slide(prs)
    _section_kicker(s, "Section 8")
    _title(s, "Commercial arrangement")
    lead = _add_textbox(s, PptInches(0.7), PptInches(1.45), PptInches(11.9), PptInches(1.0))
    lead.text_frame.word_wrap = True
    _set_para(
        lead.text_frame.paragraphs[0],
        "We welcome a good-faith conversation about an appropriate commercial arrangement associated with permission to include applicable repertoire — so parishes can prepare complete Mass materials while respecting the Publisher’s rights.",
        size=15,
        color=INK,
        space_after=0,
    )
    factors = [
        ("Repertoire", "Which titles / catalogs apply"),
        ("Intended use", "Mass presentation workflow"),
        ("Territory", TBA),
        ("Users", "Number / type of churches"),
        ("Usage model", "How content is used"),
    ]
    for i, (h, d) in enumerate(factors):
        left = PptInches(0.7 + i * 2.45)
        _add_round_rect(s, left, PptInches(2.7), PptInches(2.3), PptInches(1.85), SOFT)
        hb = _add_textbox(s, left + PptInches(0.15), PptInches(2.85), PptInches(2.0), PptInches(0.4))
        _set_para(hb.text_frame.paragraphs[0], h, size=13, bold=True, color=BERRY, space_after=0)
        db = _add_textbox(s, left + PptInches(0.15), PptInches(3.35), PptInches(2.0), PptInches(1.0))
        db.text_frame.word_wrap = True
        _set_para(db.text_frame.paragraphs[0], d, size=12, color=INK, space_after=0)
    close = _add_textbox(s, PptInches(0.7), PptInches(4.85), PptInches(11.9), PptInches(1.5))
    close.text_frame.word_wrap = True
    _set_para(
        close.text_frame.paragraphs[0],
        "Publisher permission arrangements are separate from LiturgyFlow’s public parish subscription pricing.",
        size=14,
        bold=True,
        color=INK,
        space_after=8,
    )
    _set_para(
        close.text_frame.add_paragraph(),
        "This MOA does not establish any fee, royalty, or payment obligation.",
        size=14,
        bold=True,
        color=BERRY,
        space_after=6,
    )
    _set_para(
        close.text_frame.add_paragraph(),
        f"Specific commercial terms (including any fees, royalties, or other consideration, if applicable): {TBA}.",
        size=13,
        color=MUTED,
        space_after=0,
    )
    _footer(s, 8)

    # 9 Reporting & Confidentiality
    s = _new_slide(prs)
    _section_kicker(s, "Sections 9–10")
    _title(s, "Reporting and confidentiality")
    _add_round_rect(s, PptInches(0.7), PptInches(1.55), PptInches(5.85), PptInches(4.6), SOFT)
    _add_round_rect(s, PptInches(6.75), PptInches(1.55), PptInches(5.85), PptInches(4.6), SOFT)
    lh = _add_textbox(s, PptInches(0.95), PptInches(1.75), PptInches(5.4), PptInches(0.4))
    _set_para(lh.text_frame.paragraphs[0], "Reporting", size=16, bold=True, color=BERRY, space_after=0)
    _bullets(
        s,
        PptInches(0.95),
        PptInches(2.3),
        PptInches(5.4),
        PptInches(3.5),
        [
            "LiturgyFlow may provide reasonable usage reports on the licensed repertoire",
            "Form and frequency: " + TBA,
            "Not intended to require LiturgyFlow’s complete financial records",
            "Financial reporting / audit rights: " + TBA,
        ],
        size=13,
    )
    rh = _add_textbox(s, PptInches(7.0), PptInches(1.75), PptInches(5.4), PptInches(0.4))
    _set_para(rh.text_frame.paragraphs[0], "Confidentiality", size=16, bold=True, color=BERRY, space_after=0)
    _bullets(
        s,
        PptInches(7.0),
        PptInches(2.3),
        PptInches(5.4),
        PptInches(3.5),
        [
            "Covers non-public business, technical, and commercial information exchanged",
            "Use only for evaluating and negotiating the collaboration",
            "Standard public-knowledge / independent-development carve-outs",
            "No requirement to disclose proprietary platform details beyond what LiturgyFlow chooses to share",
            "Duration and remedies: " + TBA,
        ],
        size=13,
    )
    _footer(s, 9)

    # 10 Term & Limitations
    s = _new_slide(prs)
    _section_kicker(s, "Sections 11–12")
    _title(s, "Term, termination, and limitations")
    _add_round_rect(s, PptInches(0.7), PptInches(1.55), PptInches(11.9), PptInches(1.8), SOFT)
    tbox = _add_textbox(s, PptInches(0.95), PptInches(1.75), PptInches(11.4), PptInches(1.4))
    tbox.text_frame.word_wrap = True
    _set_para(tbox.text_frame.paragraphs[0], "Term and termination", size=14, bold=True, color=BERRY, space_after=6)
    _set_para(
        tbox.text_frame.add_paragraph(),
        f"Duration, renewal, and termination terms: {TBA}. Until a definitive agreement is signed, either Party may discontinue discussions at any time.",
        size=15,
        color=INK,
        space_after=0,
    )
    lim = _add_textbox(s, PptInches(0.7), PptInches(3.7), PptInches(11.9), PptInches(0.4))
    _set_para(lim.text_frame.paragraphs[0], "Limitations", size=14, bold=True, color=BERRY, space_after=0)
    _bullets(
        s,
        PptInches(0.7),
        PptInches(4.2),
        PptInches(11.9),
        PptInches(2.0),
        [
            "Applies only to the repertoire and uses expressly agreed by the Parties",
            "No broader rights are implied",
            f"Governing law, dispute resolution, indemnity, liability, and warranties: {TBA}",
        ],
        size=15,
    )
    _footer(s, 10)

    # 11 Signatures
    s = _new_slide(prs)
    _section_kicker(s, "Section 13")
    _title(s, "Signatures — discussion acknowledgment")
    note = _add_textbox(s, PptInches(0.7), PptInches(1.45), PptInches(11.9), PptInches(1.0))
    note.text_frame.word_wrap = True
    _set_para(
        note.text_frame.paragraphs[0],
        "Signatures confirm willingness to continue good-faith discussions. They do not, by themselves, create a binding license or commercial payment obligation unless a definitive written agreement expressly states otherwise.",
        size=14,
        color=MUTED,
        space_after=0,
    )
    for i, party in enumerate(("FOR LITURGYFLOW", "FOR JESUIT MUSIC MINISTRY (JMM)")):
        left = PptInches(0.7 + i * 6.2)
        _add_round_rect(s, left, PptInches(2.7), PptInches(5.9), PptInches(3.4), SOFT)
        hb = _add_textbox(s, left + PptInches(0.3), PptInches(2.9), PptInches(5.3), PptInches(0.4))
        _set_para(hb.text_frame.paragraphs[0], party, size=13, bold=True, color=BERRY, space_after=0)
        lines = _add_textbox(s, left + PptInches(0.3), PptInches(3.5), PptInches(5.3), PptInches(2.3))
        lines.text_frame.word_wrap = True
        for j, label in enumerate(("Authorized signature: _______________________", "Name: _____________________________________", "Title: _____________________________________", "Date: _____________________________________")):
            p = lines.text_frame.paragraphs[0] if j == 0 else lines.text_frame.add_paragraph()
            run = p.add_run()
            run.text = label
            _set_run_font(run, size_pt=13, color=INK)
            p.space_after = PptPt(14)
    _footer(s, 11)

    # 12 Closing
    s = _new_slide(prs)
    _add_rect(s, 0, 0, SLIDE_W, SLIDE_H, SOFT)
    _add_rect(s, 0, 0, PptInches(0.18), SLIDE_H, BERRY)
    banner = _add_textbox(s, PptInches(0.85), PptInches(2.2), PptInches(11), PptInches(0.4))
    _set_para(banner.text_frame.paragraphs[0], "DRAFT — FOR DISCUSSION PURPOSES ONLY", size=14, bold=True, color=BERRY, space_after=0)
    title = _add_textbox(s, PptInches(0.85), PptInches(2.7), PptInches(11), PptInches(0.7))
    _set_para(title.text_frame.paragraphs[0], "Next step: mutual review", size=32, bold=True, color=INK, space_after=0)
    body = _add_textbox(s, PptInches(0.85), PptInches(3.6), PptInches(11), PptInches(2.0))
    body.text_frame.word_wrap = True
    _set_para(
        body.text_frame.paragraphs[0],
        "This draft is subject to review and revision by the parties and their respective legal advisers.",
        size=16,
        color=INK,
        space_after=10,
    )
    _set_para(
        body.text_frame.add_paragraph(),
        "Companion full-text document: LiturgyFlow_Publisher_MOA_Draft.docx",
        size=14,
        color=MUTED,
        space_after=8,
    )
    _set_para(
        body.text_frame.add_paragraph(),
        "End of Draft MOA presentation",
        size=13,
        bold=True,
        color=BERRY,
        space_after=0,
    )
    _footer(s, 12)

    prs.save(PPTX_PATH)
    return PPTX_PATH


def _find_soffice() -> str | None:
    candidates = [
        shutil.which("soffice"),
        "/opt/homebrew/bin/soffice",
        "/usr/local/bin/soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    for path in candidates:
        if path and Path(path).is_file():
            return path
    return None


def _export_pdf(source: Path, dest: Path, *, filter_name: str) -> Path:
    """Convert source Office file to PDF via LibreOffice into dest path."""
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError(
            "LibreOffice (soffice) not found. Install LibreOffice to export PDFs, "
            "or open the .docx/.pptx and export PDF manually."
        )
    with tempfile.TemporaryDirectory(prefix="moa_pdf_") as tmp:
        tmp_dir = Path(tmp)
        cmd = [
            soffice,
            "--headless",
            "--norestore",
            "--nofirststartwizard",
            "--convert-to",
            f"pdf:{filter_name}",
            "--outdir",
            str(tmp_dir),
            str(source),
        ]
        subprocess.run(cmd, check=True, capture_output=True, text=True)
        produced = tmp_dir / f"{source.stem}.pdf"
        if not produced.is_file():
            pdfs = list(tmp_dir.glob("*.pdf"))
            if not pdfs:
                raise RuntimeError(f"LibreOffice produced no PDF for {source.name}")
            produced = pdfs[0]
        shutil.copy2(produced, dest)
    return dest


def build_pdfs() -> tuple[Path, Path]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc_pdf = _export_pdf(DOCX_PATH, PDF_DOC_PATH, filter_name="writer_pdf_Export")
    ppt_pdf = _export_pdf(PPTX_PATH, PDF_PPT_PATH, filter_name="impress_pdf_Export")
    return doc_pdf, ppt_pdf


def main() -> None:
    docx_path = build_docx()
    pptx_path = build_pptx()
    print(f"Wrote {docx_path}")
    print(f"Wrote {pptx_path}")
    try:
        doc_pdf, ppt_pdf = build_pdfs()
        print(f"Wrote {doc_pdf}")
        print(f"Wrote {ppt_pdf}")
    except Exception as exc:
        print(f"PDF export skipped: {exc}")


if __name__ == "__main__":
    main()
